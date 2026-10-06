"""D88 in RE: a workbook whose substrate is found neither in indices.xlsx nor among the built-in substrates is not loaded.

The loader logged a warning and went on: the workbook counted as loaded, with no substrate material, and the first
evaluation failed later on `mats_nk["Substrate"]` (KeyError), far from the cause. The load now stops on it, and the
log says which substrate and where it was looked for.
"""

from __future__ import annotations

from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
SAMPLE = ROOT / "example" / "example_RE" / "reverse_sample.xlsx"


def _with_substrate(tmp_path: Path, name: str) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(SAMPLE)
    ws = wb["design"]
    assert ws.cell(row=1, column=3).value == "silicon"  # the substrate cell of the design sheet
    ws.cell(row=1, column=3).value = name
    out = tmp_path / f"re_{name}.xlsx"
    wb.save(out)
    return str(out)


@pytest.fixture
def re_window(qapp, monkeypatch):
    monkeypatch.setenv("CERTUS_RE_HEADLESS", "1")
    from CERTUS_RE import CertusREApp

    win = CertusREApp()
    yield win
    win.close()


def test_a_workbook_whose_substrate_is_unknown_is_not_loaded(re_window, tmp_path):
    assert re_window.load_reverse_engineering_from_path(_with_substrate(tmp_path, "Unobtainium")) is False

    assert not getattr(re_window, "_re_loaded", False)


def test_the_same_workbook_with_a_known_substrate_loads(re_window, tmp_path):
    assert re_window.load_reverse_engineering_from_path(_with_substrate(tmp_path, "silicon")) is True

    assert re_window._re_loaded
    assert re_window._re_tabular_Sub is not None
