"""Import a folder of Photon RT measurements into the deposit database of CERTUS (example/database_depots/).

    python scripts\\importer_base_depots.py <source folder> [--sortie example\\database_depots] [--date 2026-10-11]

👤 2026-10-11: *« rajoute une database de dépôts déjà réalisés pour détermination d'indice ou RE »*, from the single
layers measured on the Essent Optics Photon RT (`Fichiers_PhotonRT/Certain_Monocouche` of the standalone CERTUS INDEX
web application). This script is the way in: it can be run again on a new folder.

WHAT IT WRITES
- `bruts/`: every measurement file as it came, untouched (`.xls`, `.xlsx`, `.csv`, `.res`).
- `mesures/<id>.csv`: one file per sample, `wavelength_nm,T[,R]` in percent, wavelengths ascending, the format INDEX
  reads (`load_spectrum_columns`, then `analyze_loaded_data` finds T and R by their headers).
- `catalogue.json`: one entry per sample, with what the instrument recorded (mode, angle, polarisation, date, range,
  pitch, the instrument settings of the Parameters sheet) and what the file NAME says of the material and substrate.

WHAT IT DOES NOT DO
- It infers nothing from the spectrum. A material or a substrate is written only when the file or sample name names
  it, with `"source": "nom"`; otherwise it is null. Nothing is guessed from an index.
- It copies no comment of the web application: only the measurement files and their own parameters.

The Photon RT workbooks (`.xls` and `.xlsx`) have two sheets: `Parameters` (one column per measurement: name, dates,
mode, angle, polarisation, settings) and `Measurement` (the wavelengths, then one column per measurement). The `.csv`
and `.res` files of the same name are exports of the same data: they are kept as raw files, not read twice.

Needs `xlrd` for the `.xls` workbooks (not a CERTUS dependency: only this importer reads them).
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sys
import unicodedata
from pathlib import Path
from typing import Any

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[1]
WORKBOOKS = (".xls", ".xlsx")
EXPORTS = (".csv", ".res")
MODES = {"transmittance": "T", "reflectance": "R"}

#: What a name may say of the material, in this order (the first match wins). Patterns on the lower-cased name.
#: "Nb" and "Ta" name the oxides deposited here (Nb2O5, Ta2O5); "si" alone stays unnamed: silicon or silica.
MATERIAUX = (
    (r"ta2o5 and sio2", "Ta2O5 / SiO2"),
    (r"zns.*yf3|yf3.*zns", "ZnS / YF3"),
    (r"sb2s3", "Sb2S3"),
    (r"\bito\b|^ito", "ITO"),
    (r"tio2", "TiO2"),
    (r"yf3", "YF3"),
    (r"zns|znvis", "ZnS"),
    (r"nb2o5|(^|[^a-z])nb|absnb", "Nb2O5"),
    (r"ta2o5|absta|(^|[^a-z])ta([^a-z]|[tr]?$)", "Ta2O5"),  # not "TabsSi": a T(abs) of Si, not tantalum
    (r"sio2|silice", "SiO2"),
)

#: What a name may say of the substrate.
SUBSTRATS = (
    (r"sur si\b", "Si"),
    (r"sur verre", "verre"),
    (r"b270", "B270"),
    (r"quartz", "quartz"),
)


def slug(text: str) -> str:
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii").lower()
    return re.sub(r"[^a-z0-9]+", "-", text).strip("-")


def from_name(text: str, table: tuple[tuple[str, str], ...]) -> str | None:
    low = text.lower()
    return next((value for pattern, value in table if re.search(pattern, low)), None)


def read_workbook(path: Path) -> tuple[list[list[Any]], list[list[Any]]]:
    """The `Parameters` and `Measurement` sheets of a Photon RT workbook, as lists of rows."""
    if path.suffix.lower() == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path), on_demand=True)
        sheets = {name: book.sheet_by_name(name) for name in book.sheet_names()}
        if not {"Parameters", "Measurement"} <= set(sheets):
            raise ValueError(f"not a Photon RT workbook: sheets {list(sheets)}")
        return ([sheets["Parameters"].row_values(i) for i in range(sheets["Parameters"].nrows)],
                [sheets["Measurement"].row_values(i) for i in range(sheets["Measurement"].nrows)])
    import openpyxl

    book = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    if not {"Parameters", "Measurement"} <= set(book.sheetnames):
        raise ValueError(f"not a Photon RT workbook: sheets {book.sheetnames}")
    return ([list(r) for r in book["Parameters"].iter_rows(values_only=True)],
            [list(r) for r in book["Measurement"].iter_rows(values_only=True)])


def number(value: Any) -> float | None:
    try:
        out = float(str(value).strip().replace(",", "."))
    except (TypeError, ValueError):
        return None
    return out if out == out else None


def channel_mode(declared: str, stage_deg: float | None, detector_deg: float | None) -> str | None:
    """T or R, read in the geometry when the workbook records it, else in the declared mode.

    In manual mode the Photon RT writes "Transmittance" whatever the detector does: on 2026-10-11, 13 columns of
    `Certain_Monocouche` declared so had the detector at 16 deg with the sample at 8 deg, a specular reflectance
    (their names end in R). A detector at 180 deg sees the transmitted beam, one at twice the stage angle the
    reflected one.
    """
    if detector_deg is not None and abs(detector_deg - 180.0) < 0.5:
        return "T"
    if detector_deg is not None and stage_deg is not None and abs(detector_deg - 2.0 * stage_deg) < 0.5:
        return "R"
    return MODES.get(declared.strip().lower())


def measurements(path: Path) -> list[dict[str, Any]]:
    """One record per measurement column: its parameters and its (wavelength, value) pairs."""
    params, data = read_workbook(path)
    table = {str(r[1]).strip(): [("" if v is None else str(v).strip()) for v in r[2:]] for r in params if len(r) > 2 and r[1]}
    names = [("" if v is None else str(v).strip()) for v in params[0][2:]]
    out = []
    for k, name in enumerate(names):
        col = k + 1
        pairs = [(number(row[0]), number(row[col])) for row in data[1:] if len(row) > col]
        pairs = [(w, v) for w, v in pairs if w is not None and v is not None]
        if not pairs:
            continue

        def param(label_start: str, k: int = k) -> str:
            key = next((key for key in table if key.lower().startswith(label_start)), None)
            values = table.get(key, []) if key else []
            return values[k] if k < len(values) else ""

        declared = param("measurement mode")
        detector = number(param("detector angle"))
        mode = channel_mode(declared, number(param("sample stage angle")), detector)
        if mode is None:
            continue
        out.append({
            "echantillon": name or path.stem,
            "mode": mode,
            "mode_declare": declared or None,
            "detecteur_deg": detector,
            "angle_deg": number(param("sample stage angle")),
            "polarisation": param("polarization") or None,
            "debut": param("start date") or None,
            "parametres": {label: param(label.lower()) for label in (
                "Averaging count", "Sampling pitch, nm", "Slit width, µm", "Spot size, mm",
                "Sample thickness, mm", "Refractive index n")},
            "points": sorted(pairs),
        })
    return out


def fingerprint(points: list[tuple[float, float]]) -> str:
    return hashlib.sha256(json.dumps(points).encode("utf-8")).hexdigest()


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("source", type=Path)
    ap.add_argument("--sortie", type=Path, default=ROOT / "example" / "database_depots")
    ap.add_argument("--date", default=None, help="date d'import écrite dans le catalogue (AAAA-MM-JJ)")
    a = ap.parse_args(argv)

    files = sorted(p for p in a.source.iterdir() if p.is_file() and p.suffix.lower() in WORKBOOKS + EXPORTS)
    stems_with_workbook = {p.stem for p in files if p.suffix.lower() in WORKBOOKS}
    (a.sortie / "bruts").mkdir(parents=True, exist_ok=True)
    (a.sortie / "mesures").mkdir(parents=True, exist_ok=True)

    depots, doublons, refus, seen = [], {}, {}, {}
    for path in files:
        shutil.copyfile(path, a.sortie / "bruts" / path.name)
        if path.suffix.lower() in EXPORTS:
            if path.stem in stems_with_workbook:
                doublons.setdefault(path.stem, []).append(path.name)
            else:
                refus[path.name] = "export sans classeur Photon RT du même nom : non lu"
            continue
        try:
            records = measurements(path)
        except Exception as exc:  # a file that cannot be read is reported, not skipped in silence
            refus[path.name] = f"{type(exc).__name__}: {exc}"
            continue
        by_sample: dict[str, list[dict[str, Any]]] = {}
        for rec in records:
            by_sample.setdefault(rec["echantillon"], []).append(rec)
        for sample, recs in by_sample.items():
            # One entry per sample when it has at most one T and one R on a common grid; otherwise one per channel.
            t = [r for r in recs if r["mode"] == "T"]
            rr = [r for r in recs if r["mode"] == "R"]
            same_grid = len(t) == 1 and len(rr) == 1 and [w for w, _ in t[0]["points"]] == [w for w, _ in rr[0]["points"]]
            groups = [t + rr] if (len(t) <= 1 and len(rr) <= 1 and (len(recs) == 1 or same_grid)) else [[r] for r in recs]
            for g, group in enumerate(groups):
                base = slug(path.stem) if slug(sample) in (slug(path.stem), "") else f"{slug(path.stem)}__{slug(sample)}"
                ident = base if len(groups) == 1 else f"{base}__{g + 1}"
                key = tuple(sorted((r["mode"], fingerprint(r["points"])) for r in group))
                if key in seen:
                    doublons.setdefault(seen[key], []).append(f"{path.name} [{sample}]")
                    continue
                seen[key] = ident
                grid = [w for w, _ in group[0]["points"]]
                columns = {r["mode"]: dict(r["points"]) for r in group}
                header = ["wavelength_nm"] + [m for m in ("T", "R") if m in columns]
                lines = [",".join(header)]
                for w in grid:
                    lines.append(",".join([repr(w)] + [repr(columns[m][w]) for m in header[1:]]))
                (a.sortie / "mesures" / f"{ident}.csv").write_text("\n".join(lines) + "\n", encoding="utf-8")
                name_text = f"{path.stem} {sample}"
                materiau, substrat = from_name(name_text, MATERIAUX), from_name(name_text, SUBSTRATS)
                depots.append({
                    "id": ident,
                    "fichier": f"mesures/{ident}.csv",
                    "brut": f"bruts/{path.name}",
                    "echantillon": sample,
                    "materiau": {"valeur": materiau, "source": "nom"} if materiau else None,
                    "substrat": {"valeur": substrat, "source": "nom"} if substrat else None,
                    "voies": [{k: r[k] for k in ("mode", "mode_declare", "detecteur_deg", "angle_deg", "polarisation", "debut")}
                              for r in group],
                    "lambda_nm": [grid[0], grid[-1]],
                    "points": len(grid),
                    "parametres": group[0]["parametres"],
                })
    catalogue = {
        "version": 1,
        "description": "Monocouches mesurées au spectrophotomètre Essent Optics Photon RT (Fichiers_PhotonRT/Certain_Monocouche).",
        "importe_le": a.date,
        "unites": {"lambda": "nm", "T": "%", "R": "%"},
        "depots": depots,
        "doublons": doublons,
        "refus": refus,
    }
    (a.sortie / "catalogue.json").write_text(json.dumps(catalogue, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"{len(depots)} dépôts, {sum(len(v) for v in doublons.values())} doublons, {len(refus)} refus -> {a.sortie}")
    for name, why in refus.items():
        print(f"  refus : {name} : {why}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
