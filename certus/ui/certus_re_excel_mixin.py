from __future__ import annotations
import logging
import copy
import time
import os
import functools
import multiprocessing
import traceback
from pathlib import Path
import sys
from typing import Any, List, Dict
import numpy as np
import pyqtgraph as pg

from certus.core.certus_core import certus_timestamp_display, setup_logging, CFG, create_module_environment, NUMERICAL_FAULT_EXCEPTIONS, get_resource_path, certus_timestamp_file
from certus.core.certus_core import __version__
from certus.core.version import APP_SUITE_VERSION

from certus.ui.certus_qt_widgets import (
    QAbstractItemView, QAbstractSpinBox, QApplication, QButtonGroup, QCheckBox,
    QColor, QComboBox, QDialog, QDoubleSpinBox, QFileDialog, QFont, QFrame,
    QGridLayout, QHBoxLayout, QHeaderView, QKeySequence, QLabel, QMessageBox,
    QPushButton, QRadioButton, QScrollArea, QShortcut, QSplitter, QStackedWidget,
    QStatusBar, QTableWidgetItem, QTabWidget, QTextEdit, QTimer, Qt, QVBoxLayout,
    QWidget,
)

from certus_physics import Layer, ObliqueTarget, init_thickness, calc_spectrum_front_wrapper, calc_spectrum_full_exact_wrapper

from certus.workers.certus_spectral_workers import EvalWorker, WarmupWorker

from certus.ui.certus_spectrum_eval_ui import (
    spectrum_eval_apply_axes_legend_scale, spectrum_eval_build_worker_cfg,
    spectrum_eval_on_finished_prepare_display, spectrum_eval_plot_curves,
    spectrum_eval_run_preamble, spectrum_eval_start_worker,
)

from certus.ui.certus_ui import (
    attach_excel_clipboard_context_menu, CertusBaseApp, CertusCard, CertusCollapsible,
    CertusScientificPlot, CertusStatusPill, CertusTheme, CertusThemeToggle,
    enable_file_drop, EnhancedProgressWidget, ExcelTableWidget, FlashyCard,
    get_certus_last_dir, install_standard_shortcuts, safe_ui_action,
    set_certus_last_dir, show_toast, WelcomeGuideWidget, create_flashy_grid,
    create_header_logo_widget, create_styled_button, create_styled_label,
    create_top_actions_bar, init_certus_app, set_certus_window_icon,
    install_skeleton_loader, remove_skeleton_loader, wrap_scientific_plot_with_toolbar,
    open_documentation, confirm_stop_with_timeout,
)

from certus.utils.certus_ux import build_premium_overrides
from certus.utils.certus_data import OPENPYXL_AVAILABLE
from certus.workers.certus_re_workers import REWorker
from certus.ui.certus_re_ui import CertusREResultsDialog

from certus.utils.certus_re_helpers import (
    RE_GUI_DEFAULT_BEAM_APERTURE_DEG, RE_GUI_DEFAULT_RE_QWOT_ALPHA, RE_HL_DELTA_RE_REG_SQRT_W,
    RE_OPTIM_POINTS_PER_TARGET, RE_PHASE2_FD_MAX_WORKERS, RE_PHASE2_FD_PARALLEL,
    RE_PHASE2_ONESIDED_SPLINE_FD, RE_PHASE4_APERTURE_SCAN_POINTS, RE_P4_BEAM_AP_BOUNDS_DEG,
    RE_P4_BEAM_N_KNOTS, RE_PHASE4_TRF_MAX_NFEV, re_qwot_penalty_weight_from_preset,
    RE_RE_DEADZONE_DELTA_RE_ABS, RE_RE_DEADZONE_QWOT_ABS, RE_SPEED_PRESETS,
    RE_SPLINE_NODE2_DEFAULT_NM, RE_SPLINE_N_KNOTS, RE_SUB_CAUCHY_TUBE_DELTA,
    RE_THICKNESS_SEARCH_RADIUS_PCT, _RE_CANONICAL_SHEETS, _RE_FT_COL_MAT, _RE_FT_COL_N,
    _RE_FT_COL_NUM, _RE_FT_COL_QW, _RE_FT_COL_THICK, _parse_re_rmse_combined_from_progress_message,
    _re_calc_spectrum_for_config, _re_cell_str, _re_find_measurement_wavelength_column,
    _re_header_is_wavelength_label, _re_header_looks_like_spectrum_title, _re_index_column_map,
    _re_index_split_header_and_data, _re_measurement_values_are_percent, _re_p4_ap_staircase_polyline,
    _re_p4_kwargs_from_opt_result, _re_p4_sort_knot_pairs, _re_parse_design_metadata_row,
    _re_parse_design_qwot_rows, _re_qwot_rmse_abs_delta_at_l0, _re_resolve_re_workbook_sheets,
    _re_rmse_combined_spectral_qwot, _re_rmse_oblique_weighted, _re_sort_results_best_for_table_and_apply,
    format_re_drift_log_triplet_pct, format_re_spline_knots_log, parse_re_column_header,
    re_apply_re_index_model, re_delta_qwot_per_layer, re_drift_result_log_suffix,
    re_interp_delta_knots_clamped, re_knots_wavelengths, re_n_corr_at_lambda_ref,
    re_substrate_cauchy_n_re_from_theta, TabularMaterial, ParsedREColumn
)
calc_spectrum_front = calc_spectrum_front_wrapper
calc_spectrum_full_exact = calc_spectrum_full_exact_wrapper


class CertusREExcelMixin:
    """CertusREExcelMixin for CERTUS_RE."""

    @staticmethod
    def _re_walk_up_indices_paths(start_dir: str | None, *, max_levels: int = 64) -> list[str]:
        """Common indices.xlsx locations in *start_dir* and each parent directory."""

        if not start_dir:
            return []

        d = Path(start_dir).resolve()

        out: list[str] = []

        for _ in range(max(1, int(max_levels))):
            out.append(str(d / "indices.xlsx"))

            out.append(str(d / "example" / "indices.xlsx"))

            out.append(str(d / "example" / "database_index" / "indices.xlsx"))

            parent = d.parent

            if parent == d:
                break

            d = parent

        return out

    def _re_indices_xlsx_candidate_paths(self, re_workbook_dir: str | None = None) -> list[str]:
        """Ordered search list for indices.xlsx with canonical DB-first policy."""

        cand: list[str] = []

        env = os.environ.get("CERTUS_INDICES_XLSX")

        if env:
            cand.append(str(Path(env.strip()).resolve()))

        root_app = Path(__file__).resolve().parent

        # Canonical location used by the whole CERTUS suite.

        cand.append(str(root_app / "example" / "database_index" / "indices.xlsx"))

        cand.extend(self._re_walk_up_indices_paths(re_workbook_dir))

        cand.append(str(Path(get_resource_path("indices.xlsx")).resolve()))

        cand.extend(self._re_walk_up_indices_paths(str(root_app)))

        try:
            cand.extend(self._re_walk_up_indices_paths(os.getcwd()))

        except (OSError, ValueError):
            pass

        seen: set[str] = set()

        out: list[str] = []

        for p in cand:
            ap = str(Path(p).resolve())

            if ap not in seen:
                seen.add(ap)

                out.append(ap)

        return out

    def _re_builtin_substrate_tabular(self, substrate_name: str, l0_ref: float) -> "TabularMaterial | None":
        """Tabular substrate from built-in analytical models (Sellmeier SiO2, Si table).

        Returns a TabularMaterial sampled on a fine grid [200, 6000] nm.
        Priority: SiO2 (Sellmeier Malitson) > Silicon (table).
        """

        sub = substrate_name.lower().strip()

        # --- SiO2 / fused silica / silice (Sellmeier Malitson 1965) ---
        _sio2_names = {"sio2", "silica", "fused silica", "silice", "fusedsilica", "quartz"}
        if sub in _sio2_names or sub.startswith("sio2") or sub.startswith("fused"):
            lam_um = np.linspace(0.20, 6.00, 1161)  # step 5 nm
            l2 = lam_um ** 2
            n2 = (
                1.0
                + 0.6961663 * l2 / (l2 - 0.0684043 ** 2)
                + 0.4079426 * l2 / (l2 - 0.1162414 ** 2)
                + 0.8974794 * l2 / (l2 - 9.896161 ** 2)
            )
            n_sio2 = np.sqrt(np.maximum(n2, 1.0))
            wls_nm = lam_um * 1000.0
            return TabularMaterial(
                wls_nm,
                n_sio2,
                np.zeros_like(n_sio2),
                l0_ref=float(l0_ref),
            )

        # --- Silicon (tabulated) ---
        if "silicon" in sub or sub in ("si", "si-wafer") or sub.startswith("if ") or sub.startswith("si-"):
            from certus_physics.materials_data import SI_K_DATA, SI_N_DATA, SI_WAVELENGTH_NM

            return TabularMaterial(
                np.asarray(SI_WAVELENGTH_NM, dtype=np.float64),
                np.asarray(SI_N_DATA, dtype=np.float64),
                np.asarray(SI_K_DATA, dtype=np.float64),
                l0_ref=float(l0_ref),
            )

        return None

    def _re_resolve_substrate_material(
        self, substrate_name: str, l0_ref: float, *, re_workbook_dir: str | None = None
    ) -> tuple["TabularMaterial | None", str, str, str]:
        """Resolve the RE substrate material and return (material, source, raw_name, normalized_name)."""

        raw_name = str(substrate_name or "").strip()
        sub_lower = raw_name.lower()
        sub_map = {
            "sapphire (al2o3)": "al2o3",
            "sapphire": "al2o3",
            "silicon (if)": "si",
            "silicon": "si",
            "si-wafer": "si",
            "si-substrate": "si",
            "d263t eco": "d263t",
            "silice": "sio2",
            "silica": "sio2",
            "fused silica": "sio2",
        }
        sub_norm = sub_map.get(sub_lower, sub_lower)

        # --- Priority 1: built-in analytical model (Sellmeier SiO2, Si) ---
        builtin = self._re_builtin_substrate_tabular(sub_norm, l0_ref)
        if builtin is None:
            builtin = self._re_builtin_substrate_tabular(sub_lower, l0_ref)
        if builtin is not None:
            return builtin, f"builtin/analytical/{sub_norm}", raw_name, sub_norm

        if not OPENPYXL_AVAILABLE:
            raise ImportError("The 'openpyxl' library is required to load external substrate indices from Excel.")

        import openpyxl as _opxl

        last_err: str | None = None

        for idx_path in self._re_indices_xlsx_candidate_paths(re_workbook_dir):
            if not Path(idx_path).is_file():
                continue

            try:
                wb = _opxl.load_workbook(idx_path, data_only=True)

            except NUMERICAL_FAULT_EXCEPTIONS as e:
                last_err = str(e)

                self.log(f"RE: could not open {idx_path}: {e}", "WARNING")

                continue

            matched = None

            # 1. Word-based exact match first to prevent false positives (e.g. 'si' matching 'sio2' sheets)
            for sh in wb.sheetnames:
                sh_lower = sh.lower()
                words = sh_lower.replace("-", " ").split()
                if sub_norm in words or sub_lower in words:
                    matched = sh
                    break

            # 2. Fallback to original fuzzy match logic if no word-based exact match found
            if matched is None:
                for sh in wb.sheetnames:
                    sh_lower = sh.lower()

                    if (
                        sub_norm in sh_lower
                        or sh_lower.startswith(sub_norm)
                        or sh_lower.replace("-", " ").split()[0] in sub_norm
                        or (len(sub_norm) >= 3 and len(sh_lower) >= 3 and sub_norm[:3] == sh_lower[:3])
                        or sub_lower in sh_lower
                        or sh_lower.startswith(sub_lower)
                        or sh_lower.replace("-", " ").split()[0] in sub_lower
                        or (len(sub_lower) >= 3 and len(sh_lower) >= 3 and sub_lower[:3] == sh_lower[:3])
                    ):
                        matched = sh

                        break

            if matched is None:
                self.log(
                    f"RE: substrate '{raw_name}' (normalized '{sub_norm}') not found in {idx_path!r} "
                    f"(sheets: {', '.join(wb.sheetnames[:8])}{'…' if len(wb.sheetnames) > 8 else ''}).",
                    "WARNING",
                )

                continue

            try:
                ws = wb[matched]

                rows = list(ws.iter_rows(values_only=True))

                data = [r for r in rows if r and isinstance(r[0], (int, float))]

                if not data:
                    self.log(f"RE: sheet {matched!r} in {idx_path!r} has no numeric rows.", "WARNING")

                    continue

                wls = np.array([float(r[0]) for r in data])

                n = np.array([float(r[1]) for r in data])

                k = np.array([float(r[2]) if len(r) > 2 and r[2] is not None else 0.0 for r in data])

                mat = TabularMaterial(wls, n, k, l0_ref=l0_ref)

                return mat, f"indices.xlsx/{matched}", raw_name, sub_norm

            except NUMERICAL_FAULT_EXCEPTIONS as e:
                last_err = str(e)

                self.log(f"RE: error reading substrate sheet in {idx_path}: {e}", "WARNING")

        fb = self._re_builtin_substrate_tabular(raw_name, l0_ref)
        if fb is not None:
            return fb, f"builtin/fallback/{raw_name}", raw_name, raw_name.lower()

        if last_err is not None:
            self.log(
                f"RE: no substrate match for '{raw_name}' (normalized '{sub_norm}'): {last_err}",
                "WARNING",
            )

        return None, "unresolved", raw_name, sub_norm

    def _load_re_substrate(
        self, substrate_name: str, l0_ref: float, *, re_workbook_dir: str | None = None
    ) -> "TabularMaterial | None":
        """Load substrate TabularMaterial."""

        mat, source, raw_name, sub_norm = self._re_resolve_substrate_material(
            substrate_name, l0_ref, re_workbook_dir=re_workbook_dir
        )
        if mat is not None:
            self._re_last_substrate_source_path = source
            self._re_last_substrate_sheet = f"analytical/{sub_norm}" if source.startswith("builtin/") else source
            self.log(
                f"RE: substrate raw='{raw_name}' normalized='{sub_norm}' source={source} n@{l0_ref:.0f}nm={mat.n4:.4f}",
                "INFO",
            )
            return mat

        return None

    def _parse_re_design(self, ws) -> tuple:
        """Parse 'design' sheet -> (lambda_ref_nm, substrate_name, qwot_list).

        Tolerant layout: lambda_ref and substrate anywhere on row 1; QWOT in the

        longest numeric column below (skips 1,2,3... row-index columns).

        """

        rows = list(ws.iter_rows(values_only=True))

        if not rows:
            raise ValueError("'design' sheet is empty.")

        lambda_ref, substrate_name = _re_parse_design_metadata_row(tuple(rows[0]))

        qwot_list = _re_parse_design_qwot_rows(rows)

        if not qwot_list:
            raise ValueError("'design' sheet contains no QWOT values.")

        return lambda_ref, substrate_name, qwot_list

    def _parse_re_index(self, ws) -> tuple:
        """Parse 'index' sheet -> (wls, n1, k1, n2, k2) as numpy arrays.

        Accepts an optional header row (text labels). Column order may vary:

        wavelength is detected by header name or defaults to column A; the next

        four data columns (left-to-right) are n1,k1,n2,k2. Missing k columns

        default to zero; missing n to 1.5.

        """

        rows = list(ws.iter_rows(values_only=True))

        if not rows:
            raise ValueError("'index' sheet is empty.")

        hdr, data_rows = _re_index_split_header_and_data(rows)

        if not data_rows:
            raise ValueError("'index' sheet has no data rows.")

        max_c = max((len(r) for r in data_rows if r), default=0)

        wl_i, i1, i2, i3, i4 = _re_index_column_map(hdr, max_c)

        compact: list[tuple] = []

        for row in data_rows:
            if not row or all(v is None for v in row):
                break

            compact.append(row)

        data_rows = compact

        n_rows = len(data_rows)

        if n_rows == 0:
            raise ValueError("'index' sheet: no numeric data found.")

        def _get(row: tuple, idx: int | None, default: float) -> float | None:

            if idx is None:
                return None

            if len(row) <= idx:
                return None

            v = row[idx]

            if v is None:
                return None

            try:
                return float(v)

            except (TypeError, ValueError):
                return None

        raw_wls: list = []

        for row in data_rows:
            raw_wls.append(_get(row, wl_i, 0.0))

        wls_arr = np.array(self._reconstruct_lambda_list(raw_wls))

        if len(wls_arr) == 0:
            raise ValueError("'index' sheet: could not reconstruct wavelength column.")

        n_rows = min(n_rows, len(wls_arr))

        def _col_arr(col_idx: int | None, default: float) -> np.ndarray:

            if col_idx is None:
                return np.full(n_rows, default, dtype=np.float64)

            out = np.empty(n_rows, dtype=np.float64)

            for i in range(n_rows):
                v = _get(data_rows[i], col_idx, default)

                out[i] = default if v is None else v

            return out

        n1 = _col_arr(i1, 1.5)

        k1 = _col_arr(i2, 0.0)

        n2 = _col_arr(i3, 1.5)

        k2 = _col_arr(i4, 0.0)

        # Keep workbook order unless the sheet headers explicitly request a swap.
        # The reverse-engineering sample relies on a stable H/L mapping from the file itself.
        return wls_arr[:n_rows], n1, k1, n2, k2

    def _parse_re_measurement(self, ws) -> tuple[np.ndarray, list[tuple[ParsedREColumn, np.ndarray]], list[str]]:
        """Parse the *measurement* sheet -> ``(wls, spectra_columns, user_warnings)``.

        ``spectra_columns``: list of ``(ParsedREColumn, ndarray)`` as **fractions**.

        The lambda column is found by header or heuristic; spectral labels via

        `parse_re_column_header` (`interpretation_notes` if uncertain).

        Values taken as **percent** if magnitudes exceed ~1.25.

        """

        rows = list(ws.iter_rows(values_only=True))

        if not rows:
            raise ValueError("'measurement' sheet is empty.")

        header = rows[0]

        if not header:
            raise ValueError("'measurement' sheet: empty header row.")

        data_rows: list[tuple] = []

        for row in rows[1:]:
            if not row or all(v is None for v in row):
                break

            data_rows.append(row)

        n_rows = len(data_rows)

        if n_rows == 0:
            raise ValueError("'measurement' sheet: no numeric data found.")

        n_header_cols = len(header)

        header_user_notes: list[str] = []

        wl_col, wl_ambiguity = _re_find_measurement_wavelength_column(tuple(header), data_rows)

        if wl_ambiguity:
            header_user_notes.append(wl_ambiguity)

            self.log(f"RE measurement: {wl_ambiguity}", "WARNING")

        col_specs: list[tuple[int, ParsedREColumn]] = []

        for i in range(n_header_cols):
            if i == wl_col:
                continue

            h = header[i] if i < len(header) else None

            if h is None:
                continue

            hs = _re_cell_str(h)

            if not hs:
                continue

            if _re_header_is_wavelength_label(hs):
                continue

            # guard: never treat a wavelength label as a spectrum column

            hsl = hs.lower()

            if "wavelength" in hsl or "length" in hsl:
                continue

            try:
                spec = parse_re_column_header(str(h))

            except ValueError as e:
                if _re_header_looks_like_spectrum_title(hs):
                    self.log(f"RE measurement: skip column {i} ({hs!r}): {e}", "WARNING")

                continue

            col_specs.append((i, spec))

            for note in spec.interpretation_notes:
                msg = f"Spectral column {i} ( {spec.raw_header} ): {note}"

                header_user_notes.append(msg)

                self.log(f"RE measurement: {msg}", "WARNING")

        if not col_specs:
            raise ValueError("'measurement' sheet: no usable spectral column (need R/T-style headers).")

        wl_lbl = _re_cell_str(header[wl_col]) if wl_col < len(header) and header[wl_col] is not None else "?"

        self.log(
            f"RE measurement: lambda column index={wl_col} ({wl_lbl!r}), {len(col_specs)} spectral channel(s).",
            "INFO",
        )

        raw_wls: list = []

        for row in data_rows:
            if len(row) > wl_col:
                raw_wls.append(row[wl_col])

            else:
                raw_wls.append(None)

        wls_arr = np.array(self._reconstruct_lambda_list(raw_wls))

        if len(wls_arr) == 0:
            raise ValueError("'measurement' sheet: could not reconstruct wavelength column.")

        n_rows = min(n_rows, len(wls_arr))

        spectra_columns: list[tuple[ParsedREColumn, np.ndarray]] = []

        for col_idx, spec in col_specs:
            raw_nums: list[float] = []

            for i in range(n_rows):
                row = data_rows[i]

                if len(row) <= col_idx:
                    raw_nums.append(np.nan)

                    continue

                v = row[col_idx]

                raw_nums.append(float(v) if isinstance(v, (int, float)) else np.nan)

            use_pct = _re_measurement_values_are_percent(raw_nums)

            out: list[float] = []

            for x in raw_nums:
                if not np.isfinite(x):
                    out.append(np.nan)

                else:
                    out.append(x / 100.0 if use_pct else x)

            if use_pct:
                self.log(
                    f"RE measurement: column {col_idx} ({spec.raw_header}) interpreted as **percent** -> scale /100.",
                    "INFO",
                )

            spectra_columns.append((spec, np.asarray(out, dtype=np.float64)))

        return (
            np.asarray(wls_arr[:n_rows], dtype=np.float64),
            spectra_columns,
            header_user_notes,
        )

    def _re_automap_three_sheet_workbook(self, wb) -> dict[str, str]:
        """Infer measurement / design / index when the workbook has exactly three sheets.

        Tries all sheet permutations (six) and keeps the first that parses without error.

        Used when titles are generic (e.g. *Sheet1*, *Feuil1*) or synonyms collide.

        """

        from itertools import permutations

        names = list(wb.sheetnames)

        if len(names) != 3:
            return {}

        keys = ("measurement", "design", "index")

        for perm in permutations(names, 3):
            trial = dict(zip(keys, perm))

            try:
                _lambda_ref, _sub, qwot = self._parse_re_design(wb[trial["design"]])

                if not qwot:
                    continue

                wls_idx, *_rest = self._parse_re_index(wb[trial["index"]])

                if len(wls_idx) < 2:
                    continue

                wls_meas, specs, _warn = self._parse_re_measurement(wb[trial["measurement"]])

                if not specs or len(wls_meas) < 2:
                    continue

            except (ValueError, TypeError, KeyError):
                continue

            return trial

        return {}

    def _re_build_load_summary_text(
        self,
        *,
        file_path: str,
        workbook_sheetnames: list[str],
        sheet_map: dict[str, str],
        lambda_ref: float,
        substrate_name: str,
        qwot_list: list[float],
        wls_idx: np.ndarray,
        wls_meas: np.ndarray,
        spectra_columns: list[tuple[ParsedREColumn, np.ndarray]],
        header_warnings: list[str],
        need_back: bool,
    ) -> str:
        """Human-readable RE load summary (metadata only, no raw spectral arrays)."""

        lines: list[str] = []

        def _push(msg: str, *, suspicious: bool = False) -> None:

            lines.append(f"!! {msg}" if suspicious else msg)

        lines.append("RE workbook interpretation summary")

        lines.append("")

        _push(f"File: {Path(file_path).resolve()}")

        _push(f"Workbook sheets ({len(workbook_sheetnames)}): {', '.join(workbook_sheetnames)}")

        _push(
            "Resolved sheets: "
            f"measurement='{sheet_map.get('measurement', '?')}', "
            f"design='{sheet_map.get('design', '?')}', "
            f"index='{sheet_map.get('index', '?')}'"
        )

        lines.append("")

        lines.append("Design")

        _push(
            f"- lambda_ref: {float(lambda_ref):.3f} nm",
            suspicious=not (200.0 <= float(lambda_ref) <= 10000.0),
        )

        _push(f"- substrate (from workbook): '{substrate_name}'")

        _push(f"- layers (QWOT entries): {len(qwot_list)}", suspicious=len(qwot_list) <= 0)

        if len(qwot_list) > 0:
            q = np.asarray(qwot_list, dtype=np.float64).ravel()

            qv = q[np.isfinite(q)]

            if qv.size > 0:
                qmin = float(np.min(qv))

                qmax = float(np.max(qv))

                _push(
                    f"- QWOT range: [{qmin:.4f}, {qmax:.4f}]",
                    suspicious=(qmin <= 0.0 or qmax > 25.0),
                )

        lines.append("")

        lines.append("Index")

        _push(f"- points: {int(len(wls_idx))}", suspicious=int(len(wls_idx)) < 20)

        if len(wls_idx) > 0:
            wi = np.asarray(wls_idx, dtype=np.float64).ravel()

            wiv = wi[np.isfinite(wi)]

            if wiv.size > 0:
                wmin_i = float(np.min(wiv))

                wmax_i = float(np.max(wiv))

                _push(
                    f"- wavelength range: [{wmin_i:.1f}, {wmax_i:.1f}] nm",
                    suspicious=(wmax_i <= wmin_i or (wmax_i - wmin_i) < 50.0),
                )

        nH = float(getattr(self._re_tabular_H, "n4", np.nan))

        nL = float(getattr(self._re_tabular_L, "n4", np.nan))

        _push(
            f"- H/L n@lambda_ref: {nH:.4f} / {nL:.4f}",
            suspicious=(not (1.0 <= nH <= 5.0) or not (1.0 <= nL <= 5.0)),
        )

        if getattr(self, "_re_tabular_Sub", None) is not None:
            sub_n4 = float(getattr(self._re_tabular_Sub, "n4", np.nan))

            src = getattr(self, "_re_last_substrate_source_path", None)

            sh = getattr(self, "_re_last_substrate_sheet", None)

            if src and sh:
                _push(
                    f"- Substrate loaded: yes ({sh} from {src}), n@lambda_ref={sub_n4:.4f}",
                    suspicious=not np.isfinite(sub_n4),
                )

            else:
                _push(
                    f"- Substrate loaded: yes, n@lambda_ref={sub_n4:.4f}",
                    suspicious=not np.isfinite(sub_n4),
                )

        else:
            _push(
                "- Substrate loaded: no (RE cannot optimize without a substrate model)",
                suspicious=True,
            )

        lines.append("")

        lines.append("Measurement")

        _push(f"- channels: {len(spectra_columns)}", suspicious=len(spectra_columns) <= 0)

        _push(f"- backside model required: {'yes' if need_back else 'no'}")

        if len(wls_meas) > 0:
            wm = np.asarray(wls_meas, dtype=np.float64).ravel()

            wmv = wm[np.isfinite(wm)]

            if wmv.size > 0:
                wmin_m = float(np.min(wmv))

                wmax_m = float(np.max(wmv))

                _push(
                    f"- wavelength range: [{wmin_m:.1f}, {wmax_m:.1f}] nm",
                    suspicious=(wmax_m <= wmin_m or (wmax_m - wmin_m) < 50.0),
                )

                npts = int(np.sum(np.isfinite(np.concatenate([v for _, v in spectra_columns]))))

                _push(f"- active target points generated: {npts}", suspicious=npts < 20)

        for i, (spec, vals) in enumerate(spectra_columns, start=1):
            n_pts = int(np.sum(np.isfinite(vals)))

            _push(
                f"  {i}. {spec.raw_header} -> {spec.target_type}, angle={spec.angle_deg:.1f}, "
                f"pol={spec.pol}, back={'on' if spec.include_backside else 'off'}, points={n_pts}",
                suspicious=(n_pts < 5),
            )

            for note in spec.interpretation_notes:
                _push(f"     note: {note}", suspicious=True)

        lines.append("")

        lines.append("Warnings / inferred assumptions")

        if header_warnings:
            for w in header_warnings:
                _push(f"- {w}", suspicious=True)

        else:
            lines.append("- none")

        if self._re_meas_lambda_min_nm is not None and self._re_meas_lambda_max_nm is not None:
            lines.append("")

            _push(
                "Fit window initialized to measurement range: "
                f"[{float(self._re_meas_lambda_min_nm):.1f}, {float(self._re_meas_lambda_max_nm):.1f}] nm",
                suspicious=(float(self._re_meas_lambda_max_nm) <= float(self._re_meas_lambda_min_nm)),
            )

        return "\n".join(lines)

    def _re_show_load_summary_dialog(self, text: str) -> None:
        """Show non-modal summary dialog after RE file load."""

        import html

        from PyQt6.QtWidgets import QTextEdit

        dlg = QDialog(self)

        dlg.setWindowTitle("RE Load Summary")

        dlg.setMinimumSize(920, 640)

        lay = QVBoxLayout(dlg)

        info = QLabel("Parsed metadata and assumptions (spectral raw arrays intentionally omitted).")

        info.setWordWrap(True)

        lay.addWidget(info)

        box = QTextEdit()

        box.setReadOnly(True)

        html_lines: list[str] = []

        for ln in text.splitlines():
            if ln.startswith("!! "):
                html_lines.append(f"<b>{html.escape(ln[3:])}</b>")

            else:
                html_lines.append(html.escape(ln))

        box.setHtml("<pre style='font-family: Consolas, monospace;'>" + "\n".join(html_lines) + "</pre>")

        lay.addWidget(box, 1)

        row = QHBoxLayout()

        btn_copy = QPushButton("Copy summary")

        btn_copy.clicked.connect(functools.partial(QApplication.clipboard().setText, text))

        btn_close = QPushButton("Close")

        btn_close.clicked.connect(dlg.close)

        row.addWidget(btn_copy)

        row.addStretch()

        row.addWidget(btn_close)

        lay.addLayout(row)

        dlg.setAttribute(Qt.WidgetAttribute.WA_DeleteOnClose, True)

        dlg.show()

    @safe_ui_action
    def load_reverse_engineering(self):
        """Load a reverse-engineering .xlsx file and configure the GUI.

        Parses three sheets (measurement / design / index), populates the

        front-table, material spinboxes and target-table, enables oblique mode,

        then activates the Run RE button.

        **Measurement**: any number of spectral columns; wavelength column is detected

        by header keywords (lambda, nm, ...) or by a monotonic numeric column. Values

        in **percent** (typical R/T > 1) are auto-scaled to fractions.

        Column titles are interpreted by :func:`parse_re_column_header` (R/T, angle,

        pol, noBK / back). If interpretation is ambiguous, a warning is logged

        and a dialog may summarize assumptions for the user.

        **Index / design** sheets tolerate column reordering (see module helpers

        ``_re_index_*`` / ``_re_parse_design_*``).

        Sheet names may use common aliases (*indices* -> *index*, etc.).  Design QWOT

        may appear as **multiple numeric columns per row** (e.g. H/L pairs), as in

        ``reverse_sample0.xlsx``.

        """

        if not OPENPYXL_AVAILABLE:
            self.log("openpyxl not installed.  Run:  pip install openpyxl", "ERROR")

            return

        f, _ = QFileDialog.getOpenFileName(self, "Load RE File", get_certus_last_dir(), "Excel (*.xlsx)")

        if not f:
            return

        self.load_reverse_engineering_from_path(f)

    def load_reverse_engineering_from_path(self, path: str) -> bool:
        """Load RE workbook from *path* (no dialog). Returns True on success."""

        if not OPENPYXL_AVAILABLE:
            self.log("openpyxl not installed.  Run:  pip install openpyxl", "ERROR")

            return False

        if not path or not Path(path).is_file():
            self.log(f"RE: file not found: {path!r}", "ERROR")

            return False

        set_certus_last_dir(path)

        try:
            self._re_loading_workbook = True
            self._re_loaded_exact_ep = None
            self._use_exact_ep = False
            _t_re = time.perf_counter()

            import openpyxl

            wb = openpyxl.load_workbook(path, data_only=True)

            self.log(f"RE: workbook opened in {time.perf_counter() - _t_re:.2f}s", "INFO")

            self._re_last_substrate_source_path = None

            self._re_last_substrate_sheet = None

            sheet_map = _re_resolve_re_workbook_sheets(list(wb.sheetnames))

            not_found = [s for s in _RE_CANONICAL_SHEETS if s not in sheet_map]

            duplicated = len(sheet_map) == 3 and len(set(sheet_map.values())) < 3

            if (not_found or duplicated) and len(wb.sheetnames) == 3:
                self.log(
                    "RE: sheet titles do not map cleanly to measurement/design/index; "
                    "trying content-based assignment (3 sheets).",
                    "WARNING",
                )

                auto_map = self._re_automap_three_sheet_workbook(wb)

                if len(auto_map) == 3:
                    sheet_map = auto_map

                    not_found = []

            if not_found:
                avail = ", ".join(repr(s) for s in wb.sheetnames)

                self.log(
                    f"RE file missing sheet(s): {', '.join(not_found)}  sheets present: {avail}",
                    "ERROR",
                )

                return False

            lambda_ref = substrate_name = qwot_list = None

            wls_idx = n1_arr = _k1 = n2_arr = _k2 = None

            wls_meas = spectra_columns = re_header_warnings = None

            parse_err: ValueError | None = None

            for _attempt in range(2):
                ws_meas = wb[sheet_map["measurement"]]

                ws_des = wb[sheet_map["design"]]

                ws_idx = wb[sheet_map["index"]]

                try:
                    _t = time.perf_counter()

                    lambda_ref, substrate_name, qwot_list = self._parse_re_design(ws_des)

                    self.log(
                        f"RE: design sheet read ({len(qwot_list)} QWOT) in {time.perf_counter() - _t:.2f}s",
                        "INFO",
                    )

                    _t = time.perf_counter()

                    wls_idx, n1_arr, _k1, n2_arr, _k2 = self._parse_re_index(ws_idx)

                    self.log(
                        f"RE: index sheet read ({len(wls_idx)} points) in {time.perf_counter() - _t:.2f}s",
                        "INFO",
                    )

                    _t = time.perf_counter()

                    wls_meas, spectra_columns, re_header_warnings = self._parse_re_measurement(ws_meas)

                    self.log(
                        f"RE: measurement sheet read ({len(wls_meas)} lambda, "
                        f"{len(spectra_columns)} channels) in "
                        f"{time.perf_counter() - _t:.2f}s",
                        "INFO",
                    )

                    parse_err = None

                    break

                except ValueError as e:
                    parse_err = e

                    if _attempt == 0 and len(wb.sheetnames) == 3:
                        auto_map = self._re_automap_three_sheet_workbook(wb)

                        if auto_map and auto_map != sheet_map:
                            self.log(
                                f"RE: parse error ({e}); retrying with content-based sheet assignment.",
                                "WARNING",
                            )

                            sheet_map = auto_map

                            continue

                    raise

            if parse_err is not None:
                raise parse_err

            assert (
                lambda_ref is not None
                and qwot_list is not None
                and wls_idx is not None
                and wls_meas is not None
                and spectra_columns is not None
                and re_header_warnings is not None
            )

            if re_header_warnings:
                try:
                    self.set_validation_status("OK")
                    for _w in re_header_warnings:
                        self.add_validation_warning(f"RE header inference: {_w}")
                except NUMERICAL_FAULT_EXCEPTIONS as exc:
                    self.logger.warning("RE validation status/warnings update skipped: %s", exc)

                if os.environ.get("CERTUS_RE_HEADLESS") == "1":
                    logging.warning(
                        "RE header warnings (headless): %s",
                        " | ".join(re_header_warnings),
                    )

                else:
                    from PyQt6.QtWidgets import QMessageBox

                    QMessageBox.warning(
                        self,
                        "RE  column header interpretation",
                        "The file was loaded, but some settings were **inferred automatically**. "
                        "Please verify they match your intent:\n\n " + "\n\n ".join(re_header_warnings),
                    )
            else:
                try:
                    self.set_validation_status("OK")
                except NUMERICAL_FAULT_EXCEPTIONS as exc:
                    self.logger.warning("RE validation status update skipped: %s", exc)

            col_desc = ", ".join(
                f"{s.raw_header}->{s.target_type} ={s.angle_deg} pol={s.pol} "
                f"back={'on' if s.include_backside else 'off'}"
                for s, _ in spectra_columns
            )

            self.log(
                f"RE: {len(qwot_list)} layers, lambda_ref={lambda_ref} nm, "
                f"substrate='{substrate_name}', columns: {col_desc}",
                "INFO",
            )

            _wls_valid = np.asarray(wls_meas, dtype=np.float64)

            _wls_valid = _wls_valid[np.isfinite(_wls_valid)]

            if _wls_valid.size > 0:
                self._re_meas_lambda_min_nm = float(np.min(_wls_valid))

                self._re_meas_lambda_max_nm = float(np.max(_wls_valid))

            else:
                self._re_meas_lambda_min_nm = None

                self._re_meas_lambda_max_nm = None

            # Build TabularMaterial objects (linear interpolation, no Cauchy)

            self._re_tabular_H = TabularMaterial(wls_idx, n1_arr, _k1, l0_ref=lambda_ref)

            self._re_tabular_L = TabularMaterial(wls_idx, n2_arr, _k2, l0_ref=lambda_ref)

            n_H_ref = self._re_tabular_H.n4

            n_L_ref = self._re_tabular_L.n4

            self.log(
                f"RE: tabular indices  H: n@{lambda_ref:.0f}nm={n_H_ref:.4f} | L: n@{lambda_ref:.0f}nm={n_L_ref:.4f}",
                "INFO",
            )

            #  Populate GUI (bulk, signals blocked)

            self._re_p4_display_beam_active = False

            self._re_p4_display_ap_knots_deg = None

            self._re_p4_display_ap_knots_nm = None

            self.blockSignals(True)

            try:
                self.l0_spin.setValue(lambda_ref)

                self._re_tabular_Sub = self._load_re_substrate(
                    substrate_name,
                    lambda_ref,
                    re_workbook_dir=str(Path(path).resolve().parent),
                )

                # Front table  alternating H / L

                self.front_table.blockSignals(True)

                self.front_table.setRowCount(0)

                for i, qwot in enumerate(qwot_list):
                    mat = "H" if i % 2 == 0 else "L"

                    self._add_front_row(mat, qwot, True)

                self.front_table.blockSignals(False)

                self._update_layer_count()

                # Enable global backside UI if any channel uses a finite-plate model

                need_back = any(spec.include_backside for spec, _ in spectra_columns)

                self.back_check.setChecked(need_back)

                if need_back:
                    self.log("RE: backside calculation enabled (at least one column).", "INFO")

                else:
                    self.log(
                        "RE: all columns are front-only / noBK  backside calculation off.",
                        "INFO",
                    )

            finally:
                self.blockSignals(False)

            # Oblique mode  angle / pol / R|T per point (table = display per channel)

            self.oblique_mode = True

            self._update_target_table_headers()

            #  Build ObliqueTarget list (stored internally, not in widget)

            step = float(wls_meas[1] - wls_meas[0]) if len(wls_meas) >= 2 else 10.0

            half = step / 2.0

            self._re_targets = []

            self._re_target_col_groups = []

            for spec, vals in spectra_columns:
                group = []

                for wl, val in zip(wls_meas, vals):
                    if np.isnan(val):
                        continue

                    t = ObliqueTarget(
                        angle=float(spec.angle_deg),
                        pol=spec.pol,
                        target_type=spec.target_type,
                        lmin=max(1.0, float(wl) - half),
                        lmax=float(wl) + half,
                        tmin=float(val),
                        tmax=float(val),
                        w=1.0,
                        on=True,
                        include_backside=spec.include_backside,
                    )

                    self._re_targets.append(t)

                    group.append(t)

                self._re_target_col_groups.append(group)

            #  Populate target table: ONE display row per spectral column

            self.target_table.blockSignals(True)

            self.target_table.setRowCount(0)

            self.target_table.blockSignals(False)

            for i_col, (spec, _) in enumerate(spectra_columns):
                self.add_target()

                r = self.target_table.rowCount() - 1

                for col, value in [
                    (1, float(spec.angle_deg)),
                ]:
                    w = self.target_table.cellWidget(r, col)

                    if w:
                        w.setValue(value)

                for col, text in [
                    (2, spec.pol),
                    (3, spec.target_type),
                ]:
                    w = self.target_table.cellWidget(r, col)

                    if w:
                        w.setCurrentText(text)

                for col, value in [
                    (4, float(wls_meas[0])),
                    (5, float(wls_meas[-1])),
                    (6, 1.0),
                ]:
                    w = self.target_table.cellWidget(r, col)

                    if w:
                        w.setValue(value)

                cw = self.target_table.cellWidget(r, 0)

                if cw:
                    chk = cw.findChild(QCheckBox)

                    if chk:
                        chk.setChecked(True)

                        # Disconnect existing slot to prevent duplicate signals/default UI behavior breaking our override

                        try:
                            chk.stateChanged.disconnect()

                        except (TypeError, RuntimeError):
                            pass

                        import functools
                        chk.stateChanged.connect(functools.partial(self._on_target_group_toggled, self._re_target_col_groups[i_col], chk))

                tip = f"{spec.raw_header}\nBackside (inconsistent plate): {'yes' if spec.include_backside else 'no'}"

                for c in range(1, 7):
                    w = self.target_table.cellWidget(r, c)

                    if w:
                        w.setToolTip(tip)

            n_pts = sum(int(np.sum(~np.isnan(v))) for _, v in spectra_columns)

            ch_summ = ", ".join(f"{int(np.sum(~np.isnan(v)))} pts ({s.raw_header})" for s, v in spectra_columns)

            self.log(
                f"RE: {n_pts} measurement points stored ({ch_summ}). 1 display row per channel.",
                "INFO",
            )

            self._re_loaded = True

            self._re_workbook_path = str(Path(path).resolve())

            if (
                self._re_meas_lambda_min_nm is not None
                and self._re_meas_lambda_max_nm is not None
                and hasattr(self, "re_fit_lambda_min_spin")
                and hasattr(self, "re_fit_lambda_max_spin")
            ):
                self.re_fit_lambda_min_spin.blockSignals(True)

                self.re_fit_lambda_max_spin.blockSignals(True)

                try:
                    self.re_fit_lambda_min_spin.setValue(float(self._re_meas_lambda_min_nm))

                    self.re_fit_lambda_max_spin.setValue(float(self._re_meas_lambda_max_nm))

                finally:
                    self.re_fit_lambda_min_spin.blockSignals(False)

                    self.re_fit_lambda_max_spin.blockSignals(False)

                self.log(
                    "RE: fit lambda window initialized from workbook measurement range "
                    f"[{float(self._re_meas_lambda_min_nm):.1f}, {float(self._re_meas_lambda_max_nm):.1f}] nm.",
                    "INFO",
                )

            self._apply_re_excel_readout_from_file(lambda_ref, spectra_columns)

            self._re_opt_a_pct = 0.0

            self._re_opt_b_pct = 0.0

            self._re_opt_f_pct = 0.0

            self._re_spline_dH = None

            self._re_spline_dL = None

            self._re_spline_lam2_nm = None

            self._re_sub_cauchy_a0 = None

            self._re_sub_cauchy_a1 = None

            self._re_sub_cauchy_a2 = None

            self._re_clear_re_nk_preview()

            self._update_qwot_from_ep(ep_current if (ep_current := getattr(self, "ep_current", None)) is not None else np.asarray([], dtype=np.float64))

            self._update_thickness_display()

            self.ep_current = np.asarray(ep_current if ep_current is not None else self.ep_current, dtype=np.float64).copy() if ep_current is not None else self.ep_current

            self._use_exact_ep = True

            self._show_initial_re_rmse()

            self.launch_re_btn.setEnabled(True)

            self.eval_btn.setEnabled(True)

            self.accumulated_evals = 0

            self._re_nfev_cumulative = 0

            self._workflow_best_rmse = float("inf")

            self._best_eval_result = None

            self._best_eval_rmse = float("inf")

            self._update_status_bar_stats()

            if os.environ.get("CERTUS_RE_HEADLESS") != "1":
                try:
                    summary_txt = self._re_build_load_summary_text(
                        file_path=path,
                        workbook_sheetnames=list(wb.sheetnames),
                        sheet_map=dict(sheet_map),
                        lambda_ref=float(lambda_ref),
                        substrate_name=str(substrate_name),
                        qwot_list=list(qwot_list),
                        wls_idx=np.asarray(wls_idx, dtype=np.float64),
                        wls_meas=np.asarray(wls_meas, dtype=np.float64),
                        spectra_columns=list(spectra_columns),
                        header_warnings=list(re_header_warnings),
                        need_back=bool(need_back),
                    )

                    self._re_show_load_summary_dialog(summary_txt)

                except (
                    ValueError,
                    TypeError,
                    RuntimeError,
                    AttributeError,
                    KeyError,
                    IndexError,
                    FileNotFoundError,
                ) as e:
                    self.log(f"RE: could not open summary dialog: {e}", "WARNING")

                self._schedule_eval(True)

                # Deferred: compute initial RMSE after run_eval finishes painting

                # (run_eval is queued via QTimer(0ms) and would overwrite the title)

                QTimer.singleShot(200, self._show_initial_re_rmse)

            self.log(
                f"RE: load finished in {time.perf_counter() - _t_re:.2f}s  click  Run RE  to run optimization.",
                "SUCCESS",
            )

            show_toast(self, f"Loaded: {Path(path).name}", "success")

            return True

        except NUMERICAL_FAULT_EXCEPTIONS :
            self.log(f"RE load error:\n{traceback.format_exc()}", "ERROR")

            return False

        finally:
            self._re_loading_workbook = False

    def _apply_re_excel_readout_from_file(
        self,
        lambda_ref_nm: float,
        spectra_columns: list,
    ) -> None:
        """Update  Excel output  labels and lock synchronized fields."""

        if not hasattr(self, "_re_readout_lambda_lbl"):
            return

        lr = float(lambda_ref_nm)

        self._re_readout_lambda_lbl.setText(
            f"<b>Reference lambda₀ (Excel design sheet output)</b> : <b>{lr:.1f} nm</b>"
        )

        specs = [s for s, _ in spectra_columns]

        n = len(specs)

        n_with = sum(1 for s in specs if s.include_backside)

        if n == 0:
            html = "<b>Substrate back face (Excel header output)</b> : no measurement columns."

        elif n_with == 0:
            html = (
                f"<b>Substrate back face (inconsistent plate  Excel output)</b> : "
                f"<b>no</b> for all <b>{n}</b> column(s) (all without back face per headers)."
            )

        elif n_with == n:
            html = (
                f"<b>Substrate back face (inconsistent plate  Excel output)</b> : "
                f"<b>yes</b> for all <b>{n}</b> column(s) (all with back face per headers)."
            )

        else:
            parts = [
                "<b>Substrate back face (Excel output)</b> : <b>mixed</b>  "
                f"{n_with} column(s) <b>with</b> back face, {n - n_with} <b>without</b> :"
            ]

            for s in specs:
                parts.append(
                    f" <i>{s.raw_header}</i> : "
                    f"{'<b>with</b> back face' if s.include_backside else '<b>without</b> back face'}"
                )

            html = "<br>".join(parts)

        self._re_backside_summary_html = html

        self._re_readout_backside_lbl.setText(html)

        if hasattr(self, "l0_spin"):
            self.l0_spin.setReadOnly(True)

            self.l0_spin.setButtonSymbols(QAbstractSpinBox.ButtonSymbols.NoButtons)

            self.l0_spin.setToolTip("Value fixed by the RE file <b>design</b> sheet  read-only.")

        if hasattr(self, "back_check"):
            self.back_check.setEnabled(False)

            self.back_check.setToolTip("State fixed by column headers on the RE <b>measurement</b> sheet  read-only.")

    def _show_initial_re_rmse(self):
        init_rmse = self._compute_re_rmse(0.0, 0.0, 0.0)
        if init_rmse is not None:
            self._update_re_spectrum_title(init_rmse, suffix="(initial)")
            self.log(f"RE initial RMSE: {init_rmse:.6f}", "INFO")

    def _re_capture_loaded_exact_ep(self) -> None:
        """Freeze the exact loaded thickness vector for the first RE RMSE."""

        try:
            if self.ep_current is None:
                return

            ep = np.asarray(self.ep_current, dtype=np.float64).ravel().copy()

            if ep.size == 0 or not np.all(np.isfinite(ep)):
                return

            if self._re_loaded_exact_ep is None or not np.array_equal(self._re_loaded_exact_ep, ep):
                self._re_loaded_exact_ep = ep

            self._use_exact_ep = True
        except Exception as exc:
            self.logger.warning("RE: failed to capture exact loaded ep: %s", exc)

    def _show_re_drifted_indices_window(
        self,
        a_pct: float = 0.0,
        b_pct: float = 0.0,
        f_pct: float = 0.0,
        *,
        spline_dH: np.ndarray | None = None,
        spline_dL: np.ndarray | None = None,
        spline_lam_node2_nm: float | None = None,
        re_envelope_scale: float = 1.0,
        cauchy_a0: float | None = None,
        cauchy_a1: float | None = None,
        cauchy_a2: float | None = None,
        parent: QWidget | None = None,
    ) -> None:
        """Table of lambda vs Re before/after correction (DeltaRe splines or legacy drift %)."""

        host = parent if parent is not None else self

        w_parent = host if isinstance(host, QWidget) else self

        tgts = self._re_filter_targets_by_fit_window(
            [t for t in getattr(self, "_re_targets", []) if getattr(t, "on", True)]
        )

        if tgts:
            wls = np.array(
                sorted({(t.lmin + t.lmax) / 2.0 for t in tgts}),
                dtype=np.float64,
            )

        else:
            wls = self._re_fallback_plot_wavelengths_nm(200)

        mats, drift_factor, use_sp = self._re_init_plot_factors(wls, spline_dH, spline_dL)

        lambda_ref = float(self.l0_spin.value())

        series: list[tuple[str, list[str], list[np.ndarray]]] = []

        if use_sp:
            _lam2 = float(spline_lam_node2_nm) if spline_lam_node2_nm is not None else float(RE_SPLINE_NODE2_DEFAULT_NM)

            _knot_wl = re_knots_wavelengths(_lam2)

            dHv = re_interp_delta_knots_clamped(
                _knot_wl,
                np.asarray(spline_dH, dtype=np.float64),
                wls,
                envelope_scale=re_envelope_scale,
            )

            dLv = re_interp_delta_knots_clamped(
                _knot_wl,
                np.asarray(spline_dL, dtype=np.float64),
                wls,
                envelope_scale=re_envelope_scale,
            )

            if mats.get("H") is not None:
                nk = np.asarray(mats["H"].get_nk(wls), dtype=np.complex128)

                re0 = np.real(nk)

                series.append(("H", ["before", "after"], [re0, re0 + dHv]))

            if mats.get("L") is not None:
                nk = np.asarray(mats["L"].get_nk(wls), dtype=np.complex128)

                re0 = np.real(nk)

                series.append(("L", ["before", "after"], [re0, re0 + dLv]))

            if mats.get("Substrate") is not None:
                nk = np.asarray(mats["Substrate"].get_nk(wls), dtype=np.complex128)

                re0 = np.real(nk)

                if cauchy_a0 is not None and cauchy_a1 is not None and cauchy_a2 is not None:
                    x_sq = (lambda_ref / wls) ** 2

                    x_qu = (lambda_ref / wls) ** 4

                    A_mat = np.column_stack([np.ones_like(wls), x_sq, x_qu])

                    c_init, _, _, _ = np.linalg.lstsq(A_mat, re0, rcond=None)

                    re_cauchy_init = c_init[0] + c_init[1] * x_sq + c_init[2] * x_qu

                    re1 = float(cauchy_a0) + float(cauchy_a1) * x_sq + float(cauchy_a2) * x_qu

                    series.append(
                        ("Substrate", ["Loaded", "Cauchy Fit", "Cauchy Optimize"], [re0, re_cauchy_init, re1])
                    )

                else:
                    series.append(("Substrate", ["Loaded"], [re0]))

        else:
            channels: list[tuple[str, str, float]] = []

            if mats.get("H") is not None:
                channels.append(("H", "H", float(a_pct)))

            if mats.get("L") is not None:
                channels.append(("L", "L", float(b_pct)))

            if mats.get("Substrate") is not None:
                channels.append(("Substrate", "Substrate", float(f_pct)))

            for disp, key, pct in channels:
                nk = np.asarray(mats[key].get_nk(wls), dtype=np.complex128)

                re0 = np.real(nk)

                if key == "Substrate" and pct == 0.0:
                    series.append((disp, ["Loade"], [re0]))

                else:
                    re1 = re0 * (1.0 + (pct / 100.0) * drift_factor)

                    series.append((disp, ["before", "after"], [re0, re1]))

        if not series:
            self.log("No H / L / Substrate material to display indices.", "WARNING")

            return

        dlg = QDialog(w_parent)

        if use_sp:
            dlg.setWindowTitle(f"Indices Re(n)  DeltaRe(H,L) splines | lambda₀={lambda_ref:.0f} nm")

        else:
            dlg.setWindowTitle(
                f"Indices Re(n)  before / after drift % | lambda₀={lambda_ref:.0f} nm | "
                f"H={a_pct:+.3f}% L={b_pct:+.3f}% sub={f_pct:+.3f}%"
            )

        set_certus_window_icon(dlg)

        dlg.resize(920, min(650, 120 + 22 * len(wls)))

        lay = QVBoxLayout(dlg)

        lay.setContentsMargins(10, 10, 10, 10)

        if use_sp:
            lay.addWidget(
                QLabel(
                    "<b>Tabulated Re(n)</b> vs <b>Re + DeltaRe(lambda)</b> (linear interpolation on knots, "
                    "clamp |DeltaRe|(lambda) envelope). Substrate: dynamic Cauchy. <b>Im(n)</b> unchanged."
                )
            )

        else:
            lay.addWidget(
                QLabel(
                    "<b>Real part of indices</b>  after = Re × (1 + <i>p</i>% × <i>t</i>3), "
                    "<i>t</i> = max(0, (lambdalambda₀)/(5200lambda₀)). <b>Im(n)</b> unchanged."
                )
            )

        headers: list[str] = ["lambda_nm"]

        for disp, subheads, arrs in series:
            for sh in subheads:
                headers.append(f"Re({disp}) {sh}")

        nrows = len(wls)

        ncols = len(headers)

        tbl = ExcelTableWidget()

        tbl.setRowCount(nrows)

        tbl.setColumnCount(ncols)

        tbl.setHorizontalHeaderLabels(headers)

        tbl.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)

        tbl.setAlternatingRowColors(True)

        tbl.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)

        for i, wl in enumerate(wls):
            tbl.setItem(i, 0, QTableWidgetItem(f"{float(wl):.4f}"))

            c = 1

            for disp, subheads, arrs in series:
                for col_idx, arr in enumerate(arrs):
                    it_a = QTableWidgetItem(f"{float(arr[i]):.6f}")

                    if col_idx > 0:
                        it_a.setForeground(QColor("#60a5fa"))

                    it_a.setTextAlignment(Qt.AlignmentFlag.AlignCenter)

                    tbl.setItem(i, c, it_a)

                    c += 1

        lay.addWidget(tbl)

        btn_row = QHBoxLayout()

        copy_b = QPushButton(" Copy (TSV)")

        def _copy_tbl():

            lines = ["\t".join(headers)]

            for r in range(tbl.rowCount()):
                row = [tbl.item(r, cc).text() if tbl.item(r, cc) else "" for cc in range(tbl.columnCount())]

                lines.append("\t".join(row))

            QApplication.clipboard().setText("\n".join(lines))

            copy_b.setText(" Copied")

            QTimer.singleShot(1500, lambda: copy_b.setText(" Copy (TSV)"))

        copy_b.clicked.connect(_copy_tbl)

        close_b = QPushButton("Close")

        close_b.clicked.connect(dlg.close)

        btn_row.addWidget(copy_b)

        btn_row.addStretch()

        btn_row.addWidget(close_b)

        lay.addLayout(btn_row)

        dlg.show()

    def export_excel(self):
        """Excel snapshot: materials, stack, targets (oblique layout when RE / oblique), last spectrum."""

        if not OPENPYXL_AVAILABLE:
            self.log("openpyxl not installed. Run:  pip install openpyxl", "ERROR")

            return

        f, _ = QFileDialog.getSaveFileName(self, "CERTUS-RE  Export snapshot", get_certus_last_dir(), "Excel (*.xlsx)")

        if not f:
            return

        set_certus_last_dir(f)
        try:
            if not getattr(self, "validation_status", None):
                self.set_validation_status("OK")
        except NUMERICAL_FAULT_EXCEPTIONS as exc:
            self.logger.warning("RE export preflight validation status skipped: %s", exc)

        try:
            import openpyxl
            wb = openpyxl.Workbook()

            ws = wb.active

            ws.title = "Configuration"

            ws.append(["CERTUS-RE", APP_SUITE_VERSION])

            ws.append([f"Generated:  {certus_timestamp_display()}"])

            ws.append([])

            ws.append(["MATERIALS"])

            _l0x = float(self.l0_spin.value())

            ws.append(["Name", f"Re(n)@lambda₀ ({_l0x:.1f} nm)"])

            for k, m in self._get_materials().items():
                ws.append([k, m.n4])

            ws.append([])

            ws.append(["FRONT STACK", f"L0 = {self.l0_spin.value()} nm"])

            ws.append(["#", "Material", "QWOT", "Thickness (nm)", "Variable"])

            ep = self.ep_current if self.ep_current is not None else []

            for i, l in enumerate(self._get_front_stack()):
                ws.append(
                    [
                        i + 1,
                        l.mat,
                        l.qwot,
                        ep[i] if i < len(ep) else 0,
                        "Yes" if l.var else "No",
                    ]
                )

            if len(ep) > 0:
                ws.append([])

                ws.append(["Total Thickness (nm)", float(np.sum(ep))])

            ws.append([])

            ws.append(["SPECTRAL TARGETS"])

            _ob_tgts = bool(self.oblique_mode) or bool(getattr(self, "_re_loaded", False))

            if _ob_tgts:
                ws.append(
                    [
                        "Active",
                        "Angle (deg)",
                        "Pol",
                        "Type",
                        "lambdamin (nm)",
                        "lambdamax (nm)",
                        "Target (R/T)",
                        "Weight",
                        "Backside",
                    ]
                )

                for t in self._get_oblique_tgts():
                    tv = 0.5 * (float(t.tmin) + float(t.tmax))

                    ws.append(
                        [
                            "Yes" if t.on else "No",
                            float(t.angle),
                            str(t.pol),
                            str(t.target_type),
                            t.lmin,
                            t.lmax,
                            tv,
                            t.w,
                            "Yes" if getattr(t, "include_backside", False) else "No",
                        ]
                    )

            else:
                ws.append(["Active", "lambdamin (nm)", "lambdamax (nm)", "Tmin", "Tmax", "Weight"])

                for t in self._get_tgts():
                    ws.append(["Yes" if t.on else "No", t.lmin, t.lmax, t.tmin, t.tmax, t.w])

            if self.last_result:
                lr = self.last_result

                spec_vis = lr.get("spectra_vis") or {}

                if lr.get("oblique_mode") and isinstance(spec_vis, dict) and len(spec_vis) > 0:
                    ws2 = wb.create_sheet("Spectrum")

                    wls = np.asarray(lr["vis"]["l"], dtype=float)

                    keys_sorted = sorted(spec_vis.keys(), key=lambda k: (float(k[0]), str(k[1])))

                    hdr: list[Any] = ["Wavelength (nm)"]

                    for ang, pol in keys_sorted:
                        hdr.append(f"R_{float(ang):g}deg_{pol}")

                        hdr.append(f"T_{float(ang):g}deg_{pol}")

                    ws2.append(hdr)

                    for i in range(int(wls.size)):
                        row: list[Any] = [float(wls[i])]

                        for ang, pol in keys_sorted:
                            ch = spec_vis[(ang, pol)]

                            row.append(float(ch["R"][i]))

                            row.append(float(ch["T"][i]))

                        ws2.append(row)

                else:
                    ws2 = wb.create_sheet("Spectrum")

                    ws2.append(["Wavelength (nm)", "Transmission"])

                    r = lr["vis"]

                    for i in range(len(r["l"])):
                        ws2.append([r["l"][i], r["Ts"][i]])

            manifest: dict[str, Any] | None = None
            try:
                from certus.core.certus_metrology import ValidationStatus
                from certus.utils.certus_services import REFitService, REFitRequest
                status_txt = str(getattr(self, "validation_status", "OK") or "OK")
                try:
                    status_val = ValidationStatus(status_txt)
                except ValueError:
                    status_val = ValidationStatus.OK
                warnings_for_manifest = list(getattr(self, "validation_warnings", []) or [])
                svc = REFitService(runner=lambda _cfg: self.last_result or {})
                seed_val = None
                seed_sources = [
                    getattr(self, "run_seed", None),
                    getattr(self, "random_seed", None),
                    getattr(self, "_loaded_config", {}).get("seed")
                    if isinstance(getattr(self, "_loaded_config", None), dict)
                    else None,
                    getattr(self, "_loaded_config", {}).get("random_seed")
                    if isinstance(getattr(self, "_loaded_config", None), dict)
                    else None,
                    getattr(self, "_loaded_config", {}).get("robustness_seed")
                    if isinstance(getattr(self, "_loaded_config", None), dict)
                    else None,
                    getattr(self, "_loaded_config", {}).get("phase_a_seed")
                    if isinstance(getattr(self, "_loaded_config", None), dict)
                    else None,
                ]
                for _seed_candidate in seed_sources:
                    if _seed_candidate is None:
                        continue
                    try:
                        seed_val = int(_seed_candidate)
                        break
                    except (TypeError, ValueError):
                        continue
                req = REFitRequest(
                    config={
                        "module": "CERTUS_RE",
                        "oblique_mode": bool(getattr(self, "oblique_mode", False)),
                        "re_loaded": bool(getattr(self, "_re_loaded", False)),
                    },
                    source_paths=[
                        p
                        for p in (
                            str(getattr(self, "_last_loaded_re_file", "") or "").strip(),
                            str(getattr(self, "_re_last_substrate_source_path", "") or "").strip(),
                        )
                        if p and p != "<builtin>"
                    ],
                    seed=seed_val,
                    app_id="CERTUS_RE",
                    app_version=__version__,
                    warnings=warnings_for_manifest,
                    status=status_val,
                )
                manifest = svc.fit(req).manifest.to_dict()
                ws_m = wb.create_sheet("Manifest")
                ws_m.append(["Key", "Value"])
                for k, v in manifest.items():
                    ws_m.append([str(k), str(v)])
            except (ValueError, TypeError, RuntimeError, AttributeError, KeyError, ImportError, OSError) as exc:
                self.logger.warning("RE manifest sheet export skipped: %s", exc)

            manifest_dict = manifest if isinstance(manifest, dict) else {}
            from certus.utils.certus_data import get_missing_manifest_fields
            missing_manifest_fields = get_missing_manifest_fields(manifest_dict)
            if missing_manifest_fields:
                self.log(
                    "Export blocked: incomplete manifest (missing: " + ", ".join(missing_manifest_fields) + ")",
                    "ERROR",
                )
                return

            wb.save(f)

            self.log(f"Exported to:  {f}", "SUCCESS")

        except NUMERICAL_FAULT_EXCEPTIONS as e:
            self.log(f"Export error: {str(e)}", "ERROR")

    def export_re_targets_vs_theory(
        self,
        ep: np.ndarray | None = None,
        a_pct: float | None = None,
        b_pct: float | None = None,
        f_pct: float | None = None,
        *,
        spline_dH: list | np.ndarray | None = None,
        spline_dL: list | np.ndarray | None = None,
        run_label: str = "best",
    ) -> None:
        """Excel export: targets vs theory (phase-2 Delta Re splines if spline_dH/L given)."""

        if not getattr(self, "_re_loaded", False):
            self.log("No RE file loaded.", "WARNING")

            return

        if ep is None:
            if self.ep_current is None:
                self.log("No thicknesses to export.", "WARNING")

                return

            ep_use = np.asarray(self.ep_current, dtype=np.float64).flatten()

        else:
            ep_use = np.asarray(ep, dtype=np.float64).flatten()

        a = float(self._re_opt_a_pct if a_pct is None else a_pct)

        b = float(self._re_opt_b_pct if b_pct is None else b_pct)

        f = float(self._re_opt_f_pct if f_pct is None else f_pct)

        stack_front = self._get_front_stack()

        qwot_sum_l0 = 0.0

        for i, layer in enumerate(stack_front):
            if i < len(ep_use):
                qwot_sum_l0 += float(self._ep_nm_to_qwot(float(ep_use[i]), layer.mat))

        sdh = np.asarray(spline_dH, dtype=np.float64).ravel() if spline_dH is not None else None

        sdl = np.asarray(spline_dL, dtype=np.float64).ravel() if spline_dL is not None else None

        lam2_ex = getattr(self, "_re_spline_lam2_nm", None)

        _esc = self._re_envelope_scale_from_gui()

        rows = self._re_build_target_theory_rows(
            ep_use,
            a,
            b,
            f,
            spline_dH=sdh,
            spline_dL=sdl,
            spline_lam_node2_nm=lam2_ex,
            re_envelope_scale=_esc,
        )

        if not rows:
            self.log("No rows to export (RE targets?).", "WARNING")

            return

        manifest_dict: dict[str, Any] = {}
        try:
            from certus.utils.certus_data import get_missing_manifest_fields
            from certus.core.certus_metrology import ValidationStatus
            from certus.utils.certus_services import REFitService, REFitRequest

            status_txt = str(getattr(self, "validation_status", "OK") or "OK")
            try:
                status_val = ValidationStatus(status_txt)
            except ValueError:
                status_val = ValidationStatus.OK
            seed_val = None
            for _seed_candidate in (
                getattr(self, "run_seed", None),
                getattr(self, "random_seed", None),
                getattr(self, "_loaded_config", {}).get("seed")
                if isinstance(getattr(self, "_loaded_config", None), dict)
                else None,
                getattr(self, "_loaded_config", {}).get("random_seed")
                if isinstance(getattr(self, "_loaded_config", None), dict)
                else None,
                getattr(self, "_loaded_config", {}).get("robustness_seed")
                if isinstance(getattr(self, "_loaded_config", None), dict)
                else None,
            ):
                if _seed_candidate is None:
                    continue
                try:
                    seed_val = int(_seed_candidate)
                    break
                except (TypeError, ValueError):
                    continue
            svc = REFitService(runner=lambda _cfg: {"rows_count": len(rows), "run_label": str(run_label)})
            req = REFitRequest(
                config={
                    "module": "CERTUS_RE",
                    "export_kind": "targets_vs_theory",
                    "rows_count": int(len(rows)),
                },
                source_paths=[
                    p
                    for p in (
                        str(getattr(self, "_last_loaded_re_file", "") or "").strip(),
                        str(getattr(self, "_re_last_substrate_source_path", "") or "").strip(),
                    )
                    if p and p != "<builtin>"
                ],
                seed=seed_val,
                app_id="CERTUS_RE",
                app_version=__version__,
                warnings=list(getattr(self, "validation_warnings", []) or []),
                status=status_val,
            )
            manifest_dict = svc.fit(req).manifest.to_dict()
            missing_manifest_fields = get_missing_manifest_fields(manifest_dict)
            if missing_manifest_fields:
                self.log(
                    "Export blocked: incomplete manifest (missing: " + ", ".join(missing_manifest_fields) + ")",
                    "ERROR",
                )
                return
        except NUMERICAL_FAULT_EXCEPTIONS as exc:
            self.log(f"Manifest generation failed: {exc}", "WARNING")
            return

        if not OPENPYXL_AVAILABLE:
            self.log("openpyxl required for Excel export.", "WARNING")

            return

        default_name = f"RE_targets_vs_theory_{run_label}_{certus_timestamp_file()}.xlsx"

        path, _ = QFileDialog.getSaveFileName(
            self,
            "Export RE  targets vs theoretical spectrum",
            str(Path(get_certus_last_dir() or ".") / default_name),
            "Excel (*.xlsx)",
        )

        if not path:
            return

        set_certus_last_dir(path)

        try:
            import openpyxl
            wb = openpyxl.Workbook()

            ws = wb.active

            ws.title = "Targets_vs_theory"

            hdr = [
                "lambda_nm",
                "theta_deg",
                "Pol",
                "R_or_T",
                "Backside",
                "Weight",
                "Target",
                "Theory",
                "Delta_theory_minus_target",
                "Sum_QWOT_at_lambda0",
                "drift_Re_H_%",
                "drift_Re_L_%",
                "drift_Re_sub_%",
            ]

            ws.append(hdr)

            for r in rows:
                ws.append(
                    [
                        round(r["lambda_nm"], 6),
                        r["angle_deg"],
                        r["pol"],
                        r["target_type"],
                        "yes" if r["include_backside"] else "no",
                        r["weight"],
                        r["target"],
                        r["theory"],
                        r["delta"],
                        round(qwot_sum_l0, 6),
                        a,
                        b,
                        f,
                    ]
                )

            meta = wb.create_sheet("Meta")

            meta.append(["CERTUS-RE export"])

            meta.append(["Generated", certus_timestamp_display()])

            meta.append(["lambda0_nm (drift reference)", float(self.l0_spin.value())])

            meta.append(["Sum_QWOT_at_lambda0 (final thicknesses, mat. n@lambda0)", qwot_sum_l0])

            _rmse_ex = self._compute_re_rmse(0.0, 0.0, 0.0)

            meta.append(["RMSE_1_over_sqrt_lambda", _rmse_ex if _rmse_ex is not None else ""])

            meta.append(["Thicknesses_nm", " ".join(f"{x:.4f}" for x in ep_use)])

            meta.append(["RE_envelope_scale", _esc])

            if sdh is not None and sdl is not None and lam2_ex is not None:
                meta.append(["RE_spline_lam_node2_nm", float(lam2_ex)])

            ws_m = wb.create_sheet("Manifest")
            ws_m.append(["Key", "Value"])
            for k, v in manifest_dict.items():
                ws_m.append([str(k), str(v)])

            wb.save(path)

            self.log(f"RE targets vs theory export: {Path(path).name}", "SUCCESS")

        except NUMERICAL_FAULT_EXCEPTIONS as ex:
            self.log(f"RE export error: {ex}", "ERROR")

            logging.exception("export_re_targets_vs_theory")
