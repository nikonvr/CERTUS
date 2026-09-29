"""DESIGN's reports record where their design came from.

The three manifests of DESIGN — the full results that end a run, « Export to Excel », the
Pareto summary — read the loaded configuration file, the validation status and warnings,
the seed and the run id on the manager (`getattr(self, "_last_config_file", "")`), which
never holds them: every report said it came from no file, with no warning and no run id,
and the Pareto summary fingerprinted an empty Pareto front. Found on 2026-09-29 by listing
every guard a class puts on a name it never holds.
"""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest

EXAMPLE = Path(__file__).resolve().parents[2] / "example" / "example_design" / "JSON-design-example.json"


def _fake_ui(config_file: Path, **extra):
    logs = []
    ui = SimpleNamespace(
        last_result={"vis": {}},
        l0_spin=SimpleNamespace(value=lambda: 550.0),
        front_table=SimpleNamespace(rowCount=lambda: 3),
        log=lambda message, level="INFO": logs.append((level, message)),
        _last_config_file=str(config_file),
        validation_status="OK",
        validation_warnings=["substrate index extrapolated"],
        _workflow_run_id="run-9",
        **extra,
    )
    return ui, logs


def test_the_full_results_manifest_names_its_configuration_and_its_run(tmp_path) -> None:
    from certus.ui.certus_design_ui_export import ExportManager

    config = tmp_path / "design.json"
    config.write_text("{}", encoding="utf-8")
    ui, _logs = _fake_ui(config)

    manifest = ExportManager(ui)._build_export_manifest()

    assert [fp["path"] for fp in manifest["input_fingerprints"]] == [str(config.resolve())]
    assert manifest["warnings"] == ["substrate index extrapolated"]
    assert manifest["params"]["run_id"] == "run-9"


def test_the_pareto_summary_names_its_configuration_and_fingerprints_its_front(tmp_path, monkeypatch) -> None:
    import certus.core.certus_core as core
    import certus.ui.certus_design_ui_plot as plot
    from certus.ui.certus_design_ui_plot import PlotManager

    monkeypatch.setattr(core, "get_resource_path", lambda name: str(tmp_path / name))
    requests, fronts = [], []
    real_request, real_service = plot.IndexFitRequest, plot.IndexFitService
    monkeypatch.setattr(plot, "IndexFitRequest", lambda **kw: requests.append(kw) or real_request(**kw))
    monkeypatch.setattr(plot, "IndexFitService", lambda runner: fronts.append(runner(None)) or real_service(runner))
    config = tmp_path / "design.json"
    config.write_text("{}", encoding="utf-8")
    record = {"best_rmse": 0.01, "dmin_rmse": 12.0, "best_mc": 0.02, "dmin_mc": 11.0, "best_fab": 0.015, "dmin_fab": 6.0}
    ui, logs = _fake_ui(config, pareto_history={4: record}, _workflow_best_rmse=0.01)

    PlotManager(ui)._export_pareto_report()

    [request] = requests
    assert request["source_paths"] == [str(config)]
    assert request["warnings"] == ["substrate index extrapolated"]
    assert fronts == [{4: record}]


@pytest.fixture(autouse=True)
def _window_ends_with_the_test(qapp, monkeypatch):
    from qt_lifecycle import qt_lifecycle

    yield from qt_lifecycle(qapp, monkeypatch, "DESIGN provenance", main_windows_only=True)


def test_the_excel_export_names_its_configuration(qapp, monkeypatch, tmp_path) -> None:
    from openpyxl import load_workbook
    from PyQt6.QtWidgets import QFileDialog

    from certus.ui.certus_design_ui import CertusDesignApp

    target = tmp_path / "design.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(target), "Excel (*.xlsx)"))
    app = CertusDesignApp()
    app.load_config(str(EXAMPLE))

    app.export_excel()

    cells = [str(c.value) for ws in load_workbook(target).worksheets for row in ws.iter_rows() for c in row if c.value]
    assert any("JSON-design-example.json" in text for text in cells)
