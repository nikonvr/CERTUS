from __future__ import annotations
from certus.ui.certus_strat_common import *

class CertusStratExportMixin:
    def _extract_stack_multipliers(self, config: dict[str, Any]) -> list[float]:
        """Return normalized stack multipliers from multiple legacy JSON shapes."""
        raw = config.get("stack_multipliers")
        if raw is None:
            raw = config.get("stack_string")
        if raw is None:
            raw = config.get("stack")
        if raw is None:
            return []
        if isinstance(raw, str):
            tokens = [t.strip() for t in raw.replace("[", "").replace("]", "").split(",") if t.strip()]
            out: list[float] = []
            for tok in tokens:
                try:
                    out.append(float(tok))
                except (TypeError, ValueError):
                    continue
            return out
        if isinstance(raw, (list, tuple)):
            out = []
            for val in raw:
                try:
                    out.append(float(val))
                except (TypeError, ValueError):
                    continue
            return out
        return []

    def _validate_strat_config(self, config: dict[str, Any]) -> list[str]:
        """Return warnings for unresolved or ambiguous STRAT config values."""
        warnings: list[str] = []
        h_raw = config.get("h_material_file")
        l_raw = config.get("l_material_file")
        sub_raw = config.get("substrate_choice")

        for label, raw, kind in (("H", h_raw, "material"), ("L", l_raw, "material"), ("substrate", sub_raw, "substrate")):
            if raw is None or str(raw).strip() == "":
                continue
            resolved = self._resolve_strat_material_name(raw, kind=kind)
            if resolved != str(raw).strip():
                warnings.append(f"{label} resolved from '{raw}' to '{resolved}'")
            widget_key = "h_material_file" if label == "H" else "l_material_file" if label == "L" else "substrate_choice"
            widget = self.widgets.get(widget_key)
            if isinstance(widget, QComboBox) and resolved and widget.findText(resolved) < 0:
                warnings.append(f"{label} value '{resolved}' is not present in the loaded combo box")

        stack = self._extract_stack_multipliers(config)
        if "stack_multipliers" in config and not stack:
            warnings.append("stack_multipliers was provided but no valid numeric multipliers could be parsed")
        if "stack_string" in config and not stack:
            warnings.append("stack_string was provided but could not be parsed into multipliers")

        return warnings

    def _validate_strat_gui_state(self, config: dict[str, Any]) -> list[str]:
        """Compare populated widget state against canonical config values."""
        warnings: list[str] = []
        checks = [
            ("h_material_file", "H"),
            ("l_material_file", "L"),
            ("substrate_choice", "substrate"),
        ]
        for key, kind in checks:
            widget = self.widgets.get(key)
            if not isinstance(widget, QComboBox):
                continue
            expected = self._resolve_strat_material_name(config.get(key), kind=kind)
            actual = widget.currentText().strip()
            if expected and actual and expected != actual:
                warnings.append(f"{key}: expected '{expected}' but combo shows '{actual}'")
        return warnings

    def _resolve_manifest_seed(self, seed_container: Any) -> int | None:
        if not isinstance(seed_container, dict):
            return None
        for _k in (
            "seed",
            "random_seed",
            "robustness_seed",
            "phase_a_seed",
            "ensemble_seed",
        ):
            _v = seed_container.get(_k)
            if _v is None:
                continue
            try:
                return int(_v)
            except (TypeError, ValueError):
                continue
        return None

    def _manifest_source_paths(self) -> list[str]:
        paths: list[str] = []
        cfg_path = str(getattr(self, "_last_config_file", "") or "").strip()
        if cfg_path:
            paths.append(cfg_path)
        try:
            db_path = str(_resolve_strat_indices_db_path() or "").strip()
        except (RuntimeError, AttributeError, TypeError, ValueError, OSError):
            db_path = ""
        if db_path:
            paths.append(db_path)
        return list(dict.fromkeys(paths))

    def on_excel_ready(self, excel_data: io.BytesIO, metadata: Dict) -> Any:
        """Handle automatic export (Excel + HTML)"""
        self.logger.debug("[DEBUG-UI] on_excel_ready entered.")

        try:
            try:
                params_for_status = metadata.get("params", {}) if isinstance(metadata, dict) else {}
                if params_for_status.get("seed") is None and params_for_status.get("random_seed") is None:
                    self.set_validation_status("WARNING_UNSEEDED_STOCHASTIC")
                    self.add_validation_warning(
                        "STRAT run uses stochastic stages without explicit seed in exported params."
                    )
                else:
                    self.set_validation_status("OK")
            except NUMERICAL_FAULT_EXCEPTIONS as exc:
                self.logger.warning("Validation status update skipped during export: %s", exc)

            report_dir = get_resource_path("reports")

            os.makedirs(report_dir, exist_ok=True)

            rmse_val = float(metadata.get("rmse", metadata.get("rmse_p95", metadata.get("rmse_mean", 0.0))))

            timestamp = certus_timestamp_file()

            try:
                src_name = ""

                if hasattr(self, "_last_config_file") and self._last_config_file:
                    src_name = "_" + Path(self._last_config_file).stem

                base_name = f"Report_STRAT{src_name}_{timestamp}_RMSE_{rmse_val:.5f}"

            except (ValueError, TypeError, AttributeError):
                base_name = f"Report_STRAT_{timestamp}_RMSE_{rmse_val:.5f}"

            excel_path = str(Path(report_dir) / f"{base_name}.xlsx")

            html_path = str(Path(report_dir) / f"{base_name}.html")

            manifest_payload_source: dict[str, Any] = self.opti_results or {}
            try:
                params_for_manifest = self.collect_params()
                status_txt = str(getattr(self, "validation_status", "OK") or "OK")
                try:
                    status_val = ValidationStatus(status_txt)
                except ValueError:
                    status_val = ValidationStatus.OK
                svc = IndexFitService(runner=lambda _cfg: manifest_payload_source)
                manifest_dict = svc.fit({
                    "config": {"module": "CERTUS_STRAT", "params": params_for_manifest},
                    "source_paths": self._manifest_source_paths(),
                    "seed": self._resolve_manifest_seed(params_for_manifest),
                    "app_id": "CERTUS_STRAT",
                    "app_version": __version__,
                    "warnings": list(getattr(self, "validation_warnings", []) or []),
                    "status": status_val.value if isinstance(status_val, ValidationStatus) else str(status_val),
                }).manifest.to_dict()
            except NUMERICAL_FAULT_EXCEPTIONS as exc:
                self.logger.warning("STRAT manifest generation failed: %s", exc)
                manifest_dict = {}

            missing_manifest_fields = get_missing_manifest_fields(manifest_dict)
            if missing_manifest_fields:
                self.logger.error(
                    "STRAT export blocked: incomplete manifest (missing: %s)",
                    ", ".join(missing_manifest_fields),
                )
                if hasattr(self, "status_label"):
                    self.status_label.setText("Export blocked: incomplete manifest")
                return

            # 1. Save Excel
            self.logger.info("[DEBUG-UI] Saving Excel report...")

            with open(excel_path, "wb") as f:
                f.write(excel_data.getvalue())
            if manifest_dict:
                try:
                    import openpyxl

                    wb_m = openpyxl.load_workbook(excel_path)
                    if "Manifest" in wb_m.sheetnames:
                        del wb_m["Manifest"]
                    ws_m = wb_m.create_sheet("Manifest")
                    ws_m.append(["Key", "Value"])
                    for k, v in manifest_dict.items():
                        ws_m.append([str(k), str(v)])
                    wb_m.save(excel_path)
                except (ImportError, OSError, ValueError, TypeError, RuntimeError) as exc:
                    self.logger.warning("STRAT manifest Excel sheet injection skipped: %s", exc)
                try:
                    manifest_path = str(Path(report_dir) / f"{base_name}.manifest.json")
                    with open(manifest_path, "w", encoding="utf-8") as mf:
                        json.dump(manifest_dict, mf, ensure_ascii=False, indent=2)
                except (OSError, ValueError, TypeError) as exc:
                    self.logger.warning("STRAT manifest JSON write skipped: %s", exc)

            self.logger.info("✅ Excel report saved: %s", excel_path)

            # 2. Generate HTML Report
            self.logger.info("[DEBUG-UI] Saving HTML report...")

            # Gather plots from GUI if available

            figures = []

            if hasattr(self, "plot_stack"):
                figures.append(self.plot_stack)

            if hasattr(self, "plot_spectrum"):
                figures.append(self.plot_spectrum)

            # Create sections

            params = metadata.get("params", {})

            sections = [
                {
                    "title": "Strategy Optimization Summary",
                    "type": "kv",
                    "content": {
                        "Strategies Found": str(metadata.get("strategies_count", 0)),
                        "Best RMSE": f"{rmse_val:.5f}",
                        "Wavelength Range": f"{params.get('scan_wl_min', 0)}-{params.get('scan_wl_max', 0)} nm",
                        "Target Layers": f"{params.get('n_target_layers', '?')}",
                    },
                }
            ]

            # Insert Methodology and Details before Summary

            methodology_sections = [
                {
                    "title": "Monitoring Methodology",
                    "type": "kv",
                    "content": {
                        "Strategy Search": "Hybrid (Dynamic Programming + Stochastic)",
                        "Monitoring Type": "Monochromatic Optical Monitoring (Single Wave/Block)",
                        "Error Compensation": "Active (Real-time Re-optimization)",
                        "Simulation Engine": "Monte Carlo (Robustness Validation)",
                    },
                },
                {
                    "title": "Simulation Details",
                    "type": "text",
                    "content": (
                        "The strategy generation uses a <strong>Hybrid Dynamic Programming</strong> approach to find the optimal layer cutting sequence. "
                        "It simulates <strong>Monochromatic Optical Monitoring</strong> (Turning/Trigger Points) with real-time error compensation, ensuring that "
                        "the designed strategy is robust against deposition errors. The final validation is performed using a "
                        "<strong>Monte Carlo</strong> engine to estimate production yield."
                    ),
                },
            ]

            all_sections = methodology_sections + sections

            if generate_html_report(html_path, "CERTUS-STRAT Report", all_sections, figures):
                self.logger.info("✅ HTML report saved: %s", html_path)
                self._reports_exported = True

            self.status_label.setText(f"✓ Saved: {base_name}")
            self.logger.info("[DEBUG-UI] on_excel_ready completed successfully.")

        except NUMERICAL_FAULT_EXCEPTIONS as e:
            self.logger.error("❌ Error during auto-export: %s", e, exc_info=True)

    def _write_export_excel(
        self,
        strats: list,
        rmse_val: float,
        excel_path: str,
        report_dir: str,
        manifest_dict: dict,
        base_name: str,
    ) -> None:
        """Build the auto-export Excel workbook (Summary + Top Strategies + Best
        Structure + Layer Extrema + optional Manifest sheet) and write it to
        ``excel_path``. Also persist the manifest as a side-car JSON next to it
        under ``report_dir``. Behavior is preserved bit-for-bit from the legacy
        inline implementation."""
        params = self.collect_params()

        t_exec = f"{self.opti_results.get('execution_time', 0):.2f}" if "execution_time" in self.opti_results else "N/A"

        n_cores = str(get_safe_worker_count())

        n_iter = "N/A"

        df_summary = pd.DataFrame(
            {
                "Parameter": [
                    "Date",
                    "High Index (H)",
                    "Low Index (L)",
                    "substrate",
                    "Target Layers",
                    "Scan Range",
                    "Nucleation WL",
                    "Execution Time (s)",
                    "Processors",
                    "Iterations",
                    "Best RMSE",
                ],
                "Value": [
                    certus_timestamp_display(),
                    params.get("nH", "N/A"),
                    params.get("nL", "N/A"),
                    params.get("substrate", "N/A"),
                    str(params.get("n_target_layers", "N/A")),
                    f"{params.get('scan_wl_min', 0)}-{params.get('scan_wl_max', 0)} nm",
                    f"{params.get('l0', 'N/A')} nm",
                    t_exec,
                    n_cores,
                    n_iter,
                    f"{rmse_val:.6f}",
                ],
            }
        )

        strategies_data = []
        for s in strats[:20]:
            if not isinstance(s, dict):
                continue
            strat = s.get("strategy", {})
            if not isinstance(strat, dict):
                continue
            strategies_data.append(
                {
                    "ID": strat.get("strategy_id", ""),
                    "RMSE": s.get("rmse", 0),
                    "Robustness": s.get("robustness_score", 0),
                    "Blocks": strat.get("n_blocks", 0),
                    "Blocks Def": str(strat.get("blocks", [])),
                }
            )

        df_strategies = pd.DataFrame(strategies_data)
        df_structure = pd.DataFrame()
        df_extrema = pd.DataFrame()

        if strats and isinstance(strats[0], dict) and isinstance(strats[0].get("strategy"), dict):
            best = strats[0]["strategy"]
            blocks = best["blocks"]

            struct_data = []
            for i, b in enumerate(blocks):
                struct_data.append(
                    {
                        "Block #": i + 1,
                        "Monitoring WL": b.get("wavelength", 0),
                        "Layer Start Index": b["start"],
                        "Layer End Index": b["end"],
                    }
                )
            df_structure = pd.DataFrame(struct_data)

            ext_data = []
            for i, dists in enumerate(best.get("extrema_distances", [])):
                layer_prof = {}
                if i < len(best.get("theoretical_layer_profile", [])):
                    layer_prof = best["theoretical_layer_profile"][i]

                def fmt(v) -> Any:
                    if v > 15.0:
                        return "not critical"
                    return f"{v:.1f}"

                extrema_items = layer_prof.get("Textrema", [])
                extrema_summary = ", ".join(
                    [
                        f"{e.get('type', '?')}@{float(e.get('d_nm', 0.0)):.1f}nm:{float(e.get('T', 0.0)) * 100:.2f}%"
                        for e in extrema_items[:8]
                    ]
                )
                if len(extrema_items) > 8:
                    extrema_summary += f", ... +{len(extrema_items) - 8}"
                ext_data.append(
                    {
                        "Layer": i + 1,
                        "Tinit (%)": f"{float(layer_prof.get('Tinit', np.nan)) * 100:.3f}" if layer_prof else "",
                        "Tfinal (%)": f"{float(layer_prof.get('Tfinal', np.nan)) * 100:.3f}" if layer_prof else "",
                        "Textrema (summary)": extrema_summary,
                        "Start - Prev Extremum (OT nm)": fmt(dists.get("prev_start", 999.0)),
                        "Start - Next Extremum (OT nm)": fmt(dists.get("next_start", 999.0)),
                        "End - Prev Extremum (OT nm)": fmt(dists.get("prev_end", 999.0)),
                        "End - Next Extremum (OT nm)": fmt(dists.get("next_end", 999.0)),
                    }
                )
            if ext_data:
                df_extrema = pd.DataFrame(ext_data)

        if OPENPYXL_AVAILABLE:
            with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
                df_summary.to_excel(writer, sheet_name="Summary", index=False)
                if not df_strategies.empty:
                    df_strategies.to_excel(writer, sheet_name="Top_Strategies", index=False)
                if not df_structure.empty:
                    df_structure.to_excel(writer, sheet_name="Best_Structure", index=False)
                if not df_extrema.empty:
                    df_extrema.to_excel(writer, sheet_name="Layer_Extrema", index=False)
                if manifest_dict:
                    pd.DataFrame([{"Key": str(k), "Value": str(v)} for k, v in manifest_dict.items()]).to_excel(
                        writer, sheet_name="Manifest", index=False
                    )
            self.logger.info(f"✅ Rich Excel saved: {Path(excel_path).name}")
        else:
            to_excel_robust(df_summary, excel_path)
            self.logger.info(f"✅ Simple Excel saved (openpyxl missing): {Path(excel_path).name}")

        if manifest_dict:
            try:
                manifest_path = str(Path(report_dir) / f"{base_name}.manifest.json")
                with open(manifest_path, "w", encoding="utf-8") as mf:
                    json.dump(manifest_dict, mf, ensure_ascii=False, indent=2)
            except (OSError, ValueError, TypeError) as exc:
                self.logger.warning("STRAT manifest JSON write skipped: %s", exc)

    def _write_export_html(
        self,
        strats: list,
        rmse_val: float,
        html_path: str,
    ) -> None:
        """Build the auto-export HTML report (overview, methodology, top
        strategies and inline figures) and write it to ``html_path``. Behavior
        is preserved bit-for-bit from the legacy inline implementation."""
        params = self.collect_params()

        sections = [
            {
                "title": "Strategy Optimization Summary",
                "type": "kv",
                "content": {
                    "Best RMSE": f"{rmse_val:.5f}",
                    "Target Layers": str(params.get("n_target_layers", "N/A")),
                    "Wavelength Range": f"{params.get('scan_wl_min', 0)}-{params.get('scan_wl_max', 0)} nm",
                    "Nucleation WL": f"{params.get('l0', 'N/A')} nm",
                    "High Index Material": params.get("nH", "N/A"),
                    "Low Index Material": params.get("nL", "N/A"),
                    "substrate": params.get("substrate", "N/A"),
                    "Execution Time": (
                        f"{self.opti_results.get('execution_time', 0):.2f} s"
                        if "execution_time" in self.opti_results
                        else "N/A"
                    ),
                    "Processors": str(get_safe_worker_count()),
                    "Iterations": "N/A",
                },
            }
        ]

        methodology_sections = [
            {
                "title": "Monitoring Methodology",
                "type": "kv",
                "content": {
                    "Strategy Search": "Hybrid (Dynamic Programming + Stochastic)",
                    "Monitoring Type": "Monochromatic Optical Monitoring (Single Wave/Block)",
                    "Error Compensation": "Active (Real-time Re-optimization)",
                    "Simulation Engine": "Monte Carlo (Robustness Validation)",
                },
            },
            {
                "title": "Simulation Details",
                "type": "text",
                "content": (
                    "The strategy generation uses a <strong>Hybrid Dynamic Programming</strong> approach to find the optimal layer cutting sequence. "
                    "It simulates <strong>Monochromatic Optical Monitoring</strong> (Turning/Trigger Points) with real-time error compensation, ensuring that "
                    "the designed strategy is robust against deposition errors. The final validation is performed using a "
                    "<strong>Monte Carlo</strong> engine to estimate production yield."
                ),
            },
        ]

        all_sections = methodology_sections + sections

        if strats:
            top_strategies = strats[:10]
            table_data = []
            for s in top_strategies:
                strat = s.get("strategy", {})
                if not isinstance(strat, dict):
                    continue
                blocks_fmt = ", ".join([f"{b['start']:.0f}-{b['end']:.0f}" for b in strat.get("blocks", [])[:3]])
                if len(strat.get("blocks", [])) > 3:
                    blocks_fmt += "..."
                table_data.append(
                    {
                        "ID": strat["strategy_id"],
                        "Blocks Count": strat.get("n_blocks", 0),
                        "Structure (nm)": blocks_fmt,
                        "RMSE Score": f"{float(s.get('rmse_p95', s.get('rmse_mean', s.get('rmse', 0.0)))):.5f}",
                        "Robustness": f"{float(s.get('robustness_score', 0.0)):.5f}",
                    }
                )
            if table_data:
                all_sections.append(
                    {
                        "title": "Top Performing Strategies",
                        "type": "table",
                        "content": table_data,
                    }
                )

        figures = []
        if hasattr(self, "plot_stack") and self.plot_stack:
            figures.append(self.plot_stack)
        if hasattr(self, "plot_spectrum") and self.plot_spectrum:
            figures.append(self.plot_spectrum)

        if generate_html_report(html_path, "CERTUS-STRAT Report", all_sections, figures):
            self.logger.info(f"✅ HTML saved: {Path(html_path).name}")

    def _auto_export_results(self) -> Any:
        """Self-export results without worker signal"""
        self.logger.info("[DEBUG-UI] _auto_export_results started.")

        if getattr(self, "_reports_exported", False):
            self.logger.info("Reports already exported via excel_ready signal - skipping duplicate self-export.")
            return

        if not self.opti_results:
            self.logger.warning("[DEBUG-UI] self.opti_results is empty/None in _auto_export_results. Aborting.")
            return

        try:
            try:
                params_for_status = self.collect_params()
                if params_for_status.get("seed") is None and params_for_status.get("random_seed") is None:
                    self.set_validation_status("WARNING_UNSEEDED_STOCHASTIC")
                    self.add_validation_warning("STRAT auto-export run uses stochastic stages without explicit seed.")
                else:
                    self.set_validation_status("OK")
            except NUMERICAL_FAULT_EXCEPTIONS as exc:
                self.logger.warning("STRAT auto-export validation status update skipped: %s", exc)

            report_dir = get_resource_path("reports")

            os.makedirs(report_dir, exist_ok=True)

            # Get RMSE from the actual best-ranked strategy/result.
            # Older payloads may keep placeholder 0.0 in the first item, so we must
            # search for the first finite positive score instead of blindly using index 0.
            # The RMSE used for export naming must stay an error metric, not a robustness score.

            strats = []

            if hasattr(self, "final_results") and isinstance(self.final_results, dict):
                strats = list(self.final_results.get("all_strategies_results", []))

            if not strats and "strategies_results" in self.opti_results:
                strats = list(self.opti_results.get("strategies_results", []))

            rmse_val = extract_best_rmse(strats)

            timestamp = certus_timestamp_file()

            try:
                src_name = ""

                if hasattr(self, "_last_config_file") and self._last_config_file:
                    src_name = "_" + Path(self._last_config_file).stem

                base_name = f"Report_STRAT{src_name}_{timestamp}_RMSE_{rmse_val:.5f}"

            except (ValueError, TypeError, AttributeError):
                base_name = f"Report_STRAT_{timestamp}_RMSE_{rmse_val:.5f}"

            excel_path = str(Path(report_dir) / f"{base_name}.xlsx")

            html_path = str(Path(report_dir) / f"{base_name}.html")

            self.logger.info("Saving reports...")

            manifest_dict: dict[str, Any] = {}
            try:
                params_for_manifest = self.collect_params()
                status_txt = str(getattr(self, "validation_status", "OK") or "OK")
                try:
                    status_val = ValidationStatus(status_txt)
                except ValueError:
                    status_val = ValidationStatus.OK
                svc = IndexFitService(runner=lambda _cfg: self.opti_results or {})
                manifest_dict = svc.fit({
                    "config": {"module": "CERTUS_STRAT", "params": params_for_manifest},
                    "source_paths": self._manifest_source_paths(),
                    "seed": self._resolve_manifest_seed(params_for_manifest),
                    "app_id": "CERTUS_STRAT",
                    "app_version": __version__,
                    "warnings": list(getattr(self, "validation_warnings", []) or []),
                    "status": status_val.value if isinstance(status_val, ValidationStatus) else str(status_val),
                }).manifest.to_dict()
            except NUMERICAL_FAULT_EXCEPTIONS as exc:
                self.logger.warning("STRAT auto-export manifest generation failed: %s", exc)
                manifest_dict = {}

            missing_manifest_fields = get_missing_manifest_fields(manifest_dict)
            if missing_manifest_fields:
                self.logger.error(
                    "STRAT auto-export blocked: incomplete manifest (missing: %s)",
                    ", ".join(missing_manifest_fields),
                )
                if hasattr(self, "status_label"):
                    self.status_label.setText("Export blocked: incomplete manifest")
                return

            # 1. EXCEL EXPORT

            try:
                self.logger.info("[DEBUG-UI] Writing Excel report...")
                self._write_export_excel(strats, rmse_val, excel_path, report_dir, manifest_dict, base_name)
                self.logger.info("[DEBUG-UI] Excel report written successfully.")

            except NUMERICAL_FAULT_EXCEPTIONS as e:
                self.logger.error(f"Excel export failed:{e}")

            # 2. HTML EXPORT

            try:
                self.logger.info("[DEBUG-UI] Writing HTML report...")
                self._write_export_html(strats, rmse_val, html_path)
                self.logger.info("[DEBUG-UI] HTML report written successfully.")

            except NUMERICAL_FAULT_EXCEPTIONS as e:
                self.logger.error(f"HTML export failed:{e}", exc_info=True)

            self.status_label.setText(f"✓ Saved: {base_name}")
            self.logger.info("[DEBUG-UI] _auto_export_results completed.")

        except NUMERICAL_FAULT_EXCEPTIONS as e:
            self.logger.error(f"❌ Self-export error:{e}", exc_info=True)

