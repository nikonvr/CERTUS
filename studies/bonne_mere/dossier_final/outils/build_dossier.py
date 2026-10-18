"""Build the closure dossier from recorded measurements without optical calls."""
from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FINAL = Path(__file__).resolve().parents[1]
STUDY = FINAL.parent
DATA = FINAL / "donnees"
FIG = FINAL / "figures"
REPORTS = STUDY / "reports"


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def fmt(value, digits=6):
    return f"{value:.{digits}f}".replace(".", ",")


def spans(values):
    result = []
    for value in values:
        if result and value == result[-1][-1] + 1:
            result[-1].append(value)
        else:
            result.append([value])
    return ", ".join(str(r[0]) if len(r) == 1 else f"{r[0]} à {r[-1]}" for r in result)


def csv_data(path):
    return np.genfromtxt(path, names=True, delimiter=",")


def chart(name, draw, height=3.5):
    fig, ax = plt.subplots(figsize=(7.1, height), layout="constrained")
    draw(fig, ax)
    path = FIG / f"{name}.png"
    fig.savefig(path, dpi=180, facecolor="white")
    plt.close(fig)
    return path


def main():
    DATA.mkdir(parents=True, exist_ok=True)
    FIG.mkdir(parents=True, exist_ok=True)
    initial = read(STUDY / "design/initial_document_QWOT_H800.json")
    design = read(STUDY / "design/retenu_96_saphir/design.json")
    selected = read(STUDY / "design/deposition_selection.json")["user_preference"]
    assert selected["common_first4_physical_wavelength_nm"] == 452
    plan_path = STUDY / selected["plan"]
    plan = read(plan_path)
    lambda_audit = read(REPORTS / "autonomie_nuit_2026-10-17/final_audit_lambda4_20261018_1126.json")
    retained = next(g for g in lambda_audit["confirmation_groups"] if g["common_wavelength_nm"] == 452)
    base = lambda_audit["confirmation_reference_pooled_statistics"]
    variants = read(REPORTS / "strat70_variantes_2026-10-18/attempt_03/summary.json")
    variant_audit = read(REPORTS / "autonomie_nuit_2026-10-17/final_audit_variants70_20261018_1111.json")
    robustness = read(REPORTS / "robustesse96_2026-10-17/summary.json")
    nominal = [130.21137747209175, 183.994693878125] + design["thickness_nm"]
    wls = [None] * 98
    for block in plan["blocks"]:
        wls[block["start"]:block["end"]] = [block["wavelength"]] * (block["end"] - block["start"])
    rates = set(plan["rate_layers"])
    assert wls[:4] == [452] * 4
    assert len(rates) == 26
    monitoring = []
    for i in range(98):
        piece = i - 1 if i >= 2 else None
        material = "H800-Nb2O5" if i % 2 == 0 else "H800-SiO2"
        monitoring.append(dict(physical_layer=i + 1, component_layer=piece,
                               role="précoating témoin" if i < 2 else "dépôt pièce et témoin",
                               material=material, nominal_thickness_nm=nominal[i],
                               mode="Rate" if i in rates else "Optique",
                               wavelength_nm=None if i in rates else wls[i]))
    component_rates = [i - 1 for i in sorted(rates)]
    component_optics = [i for i in range(1, 97) if i not in component_rates]
    assert len(component_optics) == 70
    save(DATA / "strategie_retenue_452nm.json", plan)
    save(DATA / "decision_retenue.json", selected)
    save(DATA / "design_initial_100_H800.json", initial)
    save(DATA / "design_retenu_96_saphir.json", design)
    save(DATA / "resultats_452nm.json", dict(candidate=retained, paired_reference=base,
                                           audit=lambda_audit, independent_candidate_holdout_executed=False))
    save(DATA / "resultats_variantes70.json", dict(frontier=variants["pareto_holdout"],
                                                 reference=variant_audit["reference70_pooled_statistics"]))
    shutil.copy2(STUDY / "design/retenu_96_saphir/target.csv", DATA / "cible_retenue_601points.csv")
    shutil.copy2(STUDY / "design/retenu_96_saphir/layers.csv", DATA / "recette_retenue_96couches.csv")
    shutil.copy2(STUDY / "reports/strat_precoat2_compare_2026-10-17/effective_params.json", DATA / "parametres_STRAT_effectifs.json")
    shutil.copy2(STUDY / "reports/strat_two_witnesses_2026-10-17/component_seel_calibration.json", DATA / "calibration_SEEL_piece.json")
    with (DATA / "monitoring_retenu_98operations.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(monitoring[0]), delimiter=";")
        writer.writeheader()
        writer.writerows(monitoring)
    with (DATA / "design_initial_100_H800.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream, delimiter=";")
        writer.writerow(["couche_depuis_substrat", "materiau", "QWOT_550nm", "epaisseur_H800_nm"])
        for i in range(100):
            writer.writerow([i + 1, initial["materials"][initial["material_indices"][i]], initial["QWOT"][i], initial["thickness_nm"][i]])

    # Keep the retained paired trajectories once, removing duplicate prefix/component columns.
    compressed, evidence = {}, []
    for number in [130, 131, 135, 136, 140, 141]:
        case = REPORTS / f"strat70_lambda4_2026-10-18/attempt_01/case_{number:06d}"
        record = read(case / "summary.json")
        with np.load(case / "component_scores.npz", allow_pickle=False) as arrays:
            for i in range(3):
                full = np.concatenate([arrays[f"witness_prefix_nm_{i}"][:, :2], arrays[f"thickness_nm_{i}"]], axis=1)
                assert np.array_equal(full[:, :66].view(np.uint64), np.ascontiguousarray(arrays[f"witness_prefix_nm_{i}"]).view(np.uint64))
                for key, value in [("physical98_nm", full), ("valid", arrays[f"valid_{i}"]),
                                   ("target_rmse_pp", arrays[f"target_rmse_pp_{i}"]), ("seel_nm", arrays[f"seel_nm_{i}"])]:
                    compressed[f"case{number}_noise{i}_{key}"] = value
        evidence.append(dict(evaluation=number, summary=record, original_npz_sha256=sha(case / "component_scores.npz"),
                             original_native_sha256=sha(case / "native.pkl.gz")))
    np.savez_compressed(DATA / "preuves_appariees_452nm_3graines.npz", **compressed)
    save(DATA / "preuves_appariees_452nm_index.json", dict(cases=evidence, noises_pp=[.025, .05, .1],
         representation="Physical98 thicknesses stored once; component96=columns2:; original witness66=columns:66. Float64 bits checked before compaction.",
         original_native_pickle_not_retained_after_size_cleanup=True))

    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "axes.spines.top": False,
                         "axes.spines.right": False, "axes.grid": True, "grid.alpha": .18})
    ref_spectrum = csv_data(STUDY / "design/retenu_96_saphir/reference_spectrum.csv")
    spectrum = chart("spectre_retenu", lambda fig, ax: (
        ax.plot(ref_spectrum["wavelength_nm"], ref_spectrum["target_T_percent"], color="black", lw=2, label="Cible retenue"),
        ax.plot(ref_spectrum["wavelength_nm"], ref_spectrum["T_percent"], color="#146b83", lw=1, label="Design nominal 96 saphir"),
        ax.set(xlabel="Longueur d’onde en nm", ylabel="Transmission en %", ylim=(-2, 100)), ax.legend()), 3.6)
    thickness = chart("epaisseurs_96", lambda fig, ax: (
        ax.bar(np.arange(1, 97), design["thickness_nm"], color=["#146b83" if i % 2 == 0 else "#b07b3a" for i in range(96)]),
        ax.set(xlabel="Couche pièce depuis le substrat", ylabel="Épaisseur nominale en nm")), 3.0)
    initial_fig = chart("design_initial_100", lambda fig, ax: (
        ax.bar(np.arange(1, 101), initial["QWOT"], color="#586b79"),
        ax.set(xlabel="Couche initiale depuis le substrat", ylabel="QWOT à 550 nm")), 3.1)
    lambda_rows = [json.loads(line) for line in (REPORTS / "strat70_lambda4_2026-10-18/attempt_01/evaluations.jsonl").read_text(encoding="utf-8").splitlines()]
    screens = [r for r in lambda_rows if r["stage"] == "SCREENING"]
    def lambda_plot(fig, ax):
        for zero, label, color in [(True, "Sans arrêt observé", "#146b83"), (False, "Avec arrêt observé", "#b65b43")]:
            rows = [r for r in screens if r["zero_observed_failures"] == zero and r["mean_objective_pp"] is not None]
            ax.scatter([r["common_wavelength_nm"] for r in rows], [r["mean_objective_pp"] for r in rows], s=20, label=label, c=color)
        ax.set(xlabel="Longueur d’onde commune en nm", ylabel="RMSE moyen conditionnel en points", ylim=(1.5, 8))
        ax.legend()
    lambda_fig = chart("balayage_lambda4", lambda_plot, 3.4)
    def cdf_plot(fig, ax):
        for numbers, label, color in [([130, 135, 140], "Référence 70", "#586b79"), ([131, 136, 141], "452 nm commun", "#146b83")]:
            values = np.sort(np.concatenate([compressed[f"case{n}_noise1_target_rmse_pp"] for n in numbers]))
            ax.plot(values, np.arange(1, len(values) + 1) / len(values), label=label, color=color)
        ax.set(xlabel="RMSE cible au bruit 0,05 en points", ylabel="Fraction cumulée", xlim=(1.35, 2.2))
        ax.legend()
    cdf = chart("distribution_452nm", cdf_plot, 3.2)

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.7), Inches(.65)
    section.left_margin = section.right_margin = Inches(.8)
    for name in ["Normal", "Body Text", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption", "Header", "Footer"]:
        style = doc.styles[name]
        style.font.name = "Arial" if name in {"Title", "Subtitle", "Heading 1", "Heading 2", "Caption"} else "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.1
    for name, size in [("Title", 27), ("Heading 1", 19), ("Heading 2", 13)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].paragraph_format.space_after = Pt(12)
    doc.styles["Caption"].font.size = Pt(9)
    doc.styles["Caption"].font.bold = False
    doc.styles["Caption"].font.italic = False
    doc.styles["Caption"].paragraph_format.space_after = Pt(9)
    header = section.header.paragraphs[0]
    header.text = "BONNE MÈRE   |   DOSSIER DE CLÔTURE   |   18 OCTOBRE 2026"
    header.runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Bonne Mère  •  ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    source_list = []
    def readable(text):
        text = str(text)
        if "/" in text and (".json" in text or "reports/" in text):
            return text
        protected = []
        def protect(match):
            protected.append(match.group())
            return "¤" + chr(0xE000 + len(protected) - 1) + "¤"
        text = re.sub(r"(?<![0-9a-f])[0-9a-f]{40,64}(?![0-9a-f])|\[S[1-9]\]|\bS[1-9]\b|WorkerThread23|ELITE6", protect, text)
        text = re.sub(r"(?<=\d)(?=(?:nm|points?|tirages|couches|optiques|Rate|plans|valeurs|évaluations|trajectoires|candidats|finalistes|références|sources|entrées|archives|fois|ensembles|colonnes|dérivées|premières|par|à|et|contre|de|sur|ne|est|y|sont|au|lambda|λ))", " ", text)
        text = re.sub(r"(?<=[A-Za-zÀ-ÿ])(?=\d)", " ", text)
        text = re.sub(r"(?<=\d)(?=[A-Za-zÀ-ÿ])", " ", text)
        for before, after in [("Nb 2 O 5", "Nb2O5"), ("SiO 2", "SiO2"), ("H 800", "H800"), ("id 3", "id 3"), ("95 e", "95e")]:
            text = text.replace(before, after)
        for before, after in [("de±", "de ±"), ("relatives±", "relatives ±"), ("absolues±", "absolues ±"), ("nm.", "nm."), ("deuxprécoatings", "deux précoatings"), ("deuxpremières", "deux premières"), ("sixcohortes", "six cohortes"), ("lesNPZ", "les NPZ"), ("lesprécoatings", "les précoatings"), ("lesP95", "les P95"), ("leP95", "le P95"), ("laSEEL", "la SEEL"), ("brancheSTRAT", "branche STRAT"), ("campagnesSTRAT", "campagnes STRAT"), ("branchePOEM", "branche POEM"), ("l’essaiPGLOBAL", "l’essai PGLOBAL"), ("plafondRate", "plafond Rate"), ("leWorkerThread", "le WorkerThread"), ("PhaseA", "Phase A"), ("historiquesRate", "historiques Rate"), ("coucheRate", "couche Rate"), ("uneλ", "une λ"), ("àquatreλ", "à quatre λ"), ("λcommunes", "λ communes"), ("λcommune", "λ commune"), ("λde", "λ de"), ("Valeurλ", "Valeur λ"), ("LirePASSATION", "Lire PASSATION"), ("etdecision", "et decision"), ("dossierCERTUS", "dossier CERTUS"), ("répertoireCERTUS", "répertoire CERTUS"), ("dansarchives", "dans archives"), ("dansdossier", "dans dossier"), ("fichierdesign", "fichier design"), ("fichierstrategie", "fichier strategie"), ("le70", "le 70"), ("maximumoptique", "maximum optique"), ("minimumRMSE", "minimum RMSE"), ("P95RMSE", "P95 RMSE"), ("SEELP95", "SEEL P95"), ("P95SEEL", "P95 SEEL"), ("indicesH800", "indices H800"), ("sourceMF", "source MF"), ("historiqueN=", "historique N = ")]:
            text = text.replace(before, after)
        for i, value in enumerate(protected):
            text = text.replace("¤" + chr(0xE000 + i) + "¤", value)
        return text
    def p(text, style=None):
        return doc.add_paragraph(readable(text), style)
    def page(title):
        doc.add_page_break()
        doc.add_heading(title, 1)
    def h(text):
        doc.add_heading(readable(text), 2)
    def picture(path, caption, width=6.7):
        doc.add_picture(str(path), width=Inches(width))
        p(caption, "Caption")
    def table(headers, rows, widths=None, size=10):
        t = doc.add_table(rows=1, cols=len(headers))
        t.autofit = False
        for i, text in enumerate(headers):
            t.rows[0].cells[i].text = readable(text)
        for row in rows:
            cells = t.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = readable(value)
        for ri, row in enumerate(t.rows):
            if ri == 0:
                repeat = OxmlElement("w:tblHeader")
                row._tr.get_or_add_trPr().append(repeat)
            cant = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(cant)
            for ci, cell in enumerate(row.cells):
                if widths:
                    cell.width = Inches(widths[ci])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props = cell._tc.get_or_add_tcPr()
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "DCE8ED" if ri == 0 else ("F4F6F7" if ri % 2 == 0 else "FFFFFF"))
                props.append(shading)
                margins = OxmlElement("w:tcMar")
                for side, value in [("top", "55"), ("bottom", "55"), ("left", "80"), ("right", "80")]:
                    x = OxmlElement(f"w:{side}")
                    x.set(qn("w:w"), value)
                    x.set(qn("w:type"), "dxa")
                    margins.append(x)
                props.append(margins)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.02
                    if ci > 0:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(size)
                        run.bold = ri == 0
        borders = OxmlElement("w:tblBorders")
        for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            x = OxmlElement(f"w:{side}")
            x.set(qn("w:val"), "single")
            x.set(qn("w:sz"), "4")
            x.set(qn("w:color"), "D9D9D9")
            borders.append(x)
        t._tbl.tblPr.append(borders)
        p("").paragraph_format.space_after = Pt(2)
        return t

    doc.add_heading("Bonne Mère", 0)
    p("Du design initial à la stratégie de dépôt retenue", "Subtitle")
    p("Dossier scientifique de clôture et de reprise différée\nPréparé pour Fabien le 18 octobre 2026")
    h("Décision retenue")
    p("Conserver 70 couches optiques et 26 couches en Rate sur un seul témoin saphir, avec 452 nm pour les deux précoatings du témoin et les deux premières couches de la pièce. La recette finale comporte 96 couches H800 sur saphir. Ce choix privilégie la simplification du monitoring initial et les résultats de précision les plus prometteurs de cette branche.")
    table(["Paramètre", "Valeur retenue"], [["Design", "96 couches H800 sur saphir"], ["Précoatings témoin", "130,211377 nm puis 183,994694 nm"], ["Monitoring des quatre premières opérations", "452 nm commun"], ["Modes pièce", "70 optiques et 26 Rate"], ["Statut", "Choix utilisateur après confirmation seulement"]], [3.3, 3.6])
    p("La piste452nm a donné un arrêt sur3000tirages au bruit0,10, aucun aux deux bruits inférieurs. Elle n’a donc pas franchi le critère strict zéro arrêt et n’a pas reçu de validation indépendante finale. Cette réserve fait partie du choix enregistré, sans qualification expérimentale. [S8]")
    p("Toutes les recherches sont arrêtées. Le présent dossier rassemble la décision, les preuves conservées et les conditions d’une reprise ultérieure ; il ne demande aucun calcul nouveau.")

    page("1 Origine du problème et données de départ")
    p("Le problème Bonne Mère associe la conception d’un filtre multicouche et la recherche d’une stratégie de fabrication compatible avec le matériel de monitoring. Les pièces initiales sont un document Word, un PowerPoint et un classeur Excel. La cible retrouvée par Fabien complète ces sources. [S1]")
    table(["Source", "Contenu utile"], [["BM.docx et BM.pptx", "Contexte et figures historiques"], ["data.xlsx feuille design", "100 couches et QWOT à550nm"], ["target_retrouvee.txt", "601ordonnées de transmission retrouvées"], ["Tables H800", "Indices complexes des deux matériaux déposés"]], [2.6, 4.3])
    h("Une ambiguïté spectrale résolue explicitement")
    p("Le graphique cible du PowerPoint utilise400à1000nm et5à90% de transmission. Le texte retrouvé porte un axe600à725nm et0à85%. Le spectre Excel est un calcul du design historique sur une grille quasi logarithmique400à728,629nm. Ces objets ne sont pas interchangeables.")
    p("La décision de travail a consisté à conserver les601ordonnées retrouvées et à leur attribuer les longueurs d’onde historiques du classeur. La transformation affine ultérieure de l’axe, conjointement aux épaisseurs, aboutit à400à728,849591640315nm pour le design saphir retenu. Les ordonnées restent inchangées. [S2]")
    p("Le commentaire historiqueN=78 ne décrit pas le classeur exploité :100couches y sont réellement présentes. Le mérite sourceMF=1,60435 n’a pas de définition assez complète pour être assimilé au RMSE employé ici.")
    h("Convention de numérotation")
    p("Les couches pièce sont numérotées du substrat vers l’air. Dans le plan de dépôt, les opérations physiques1et2 sont les précoatings du témoin ; la couche pièce1 correspond à l’opération3. Les précoatings ne sont pas compris dans les96couches pièce.")

    page("2 Le design initial à cent couches")
    p("Le point de départ conservé est la série des100QWOT à550nm lue directement dans la feuille design du classeur. L’initialisation n’est pas un candidat déjà optimisé. Les matériaux sont convertis vers H800-Nb2O5 et H800-SiO2 en conservant ordre, alternance et QWOT. [S3]")
    p("La conversion appliquée est : épaisseur en nm = QWOT ×550 / (4×indice réel H800 à550nm). Les indices de référence enregistrés sont2,357884 et1,486816. Les épaisseurs documentaires initiales et les épaisseurs après conversionH800 restent distinguées dans les sources.")
    picture(initial_fig, "Figure1  Les cent QWOT documentaires conservés comme point de départ")
    p("La recherchePGLOBAL native explore les100variablesQWOT sans polissage préalable, avec transformation exacte du gradient en épaisseurs physiques. Les contrôles préalables portent sur l’oracle optique et les100dérivées. Les données initiales complètes figurent en annexe et dans le CSV associé.")
    p("Les scores de départ anciens dépendent du protocole de cible, de la face arrière et de l’échantillonnage. Le présent dossier ne calcule pas un pourcentage de gain global en juxtaposant ces repères : la cible et le substrat ont évolué pendant la conception.")

    page("3 Modèle optique et critères")
    p("Les calculs scientifiques utilisent le moteur CERTUS figé au commit7f3edea879ab150eb9bee50b3b9f5d1675bb3869. Le présent dossier lit les résultats déjà enregistrés ; il ne contient aucune nouvelle simulation optique ni réimplémentation TMM. [S1]")
    table(["Élément", "Hypothèse enregistrée"], [["Matériaux déposés", "H800-Nb2O5 et H800-SiO2 avec absorption"], ["Convention", "Indice complexe n−ik avec k positif"], ["Substrat initial", "Silice massive dispersive"], ["Substrat retenu", "Saphir modèle ordinaire natif id3"], ["Face arrière", "Nue incluse dans le modèle retenu"], ["Incidence", "Normale"], ["Cible finale", "601points sur400à728,849592nm"]], [2.3, 4.6])
    h("Sens des grandeurs rapportées")
    p("Le RMSE primaire est l’écart quadratique moyen non pondéré de transmission aux601points de la cible, exprimé en points de pourcentage. Le RMSE moyen de dépôt est ensuite la moyenne de ce critère sur les trajectoires terminées. LeP95 est son95epercentile. La SEEL exprime une erreur d’épaisseur équivalente sur la pièce, en nm, avec la calibration enregistrée.")
    p("Les statistiques de précision sont conditionnelles aux dépôts terminés. Les arrêts sont toujours publiés séparément. Pour regrouper trois graines, les quantiles sont relus dans les tableaux concaténés ; les troisP95 ne sont pas simplement moyennés.")
    p("Un score sur une grille dense est un contrôle secondaire et ne remplace pas le RMSE primaire. Une comparaison de stratégies n’est recevable que sur les mêmes graines, budgets et bruits.")

    page("4 Du design initial au design retenu")
    p("La conception a exploré des architectures de100à117couches, des regroupements et suppressions de couches minces, puis un compromis à96couches. Des minima nominaux plus bas avec davantage de couches ont été conservés comme repères ; ils ne définissent pas le choix final à eux seuls. [S3,S4]")
    best100 = read(REPORTS / "qwot_global_historique_01/best.json")
    local116 = read(REPORTS / "local_116_sans_sio2_2026-10-17/summary.json")
    local96 = read(REPORTS / "local_96_sans_plancher_2026-10-17/n96/summary.json")
    table(["Étape", "Couches", "RMSE nominal", "Statut"], [["QWOT documentaires", "100", "Protocole initial", "Départ conservé"], ["PGLOBAL100 natif", "100", fmt(best100["reference_rmse_percentage_points"]), "Minimum trouvé"], ["Variante locale silice", "116", fmt(local116["local_rmse_pp"]), "Repère archivé"], ["96silice sans plancher5nm", "96", fmt(local96["local_rmse_pp"]), "Référence robuste"], ["96saphir et axe conjoint", "96", fmt(design["reference_rmse_percentage_points"]), "Design choisi"]], [2.55, .6, 1.45, 2.3], 9.4)
    p("Les lignes silice et saphir n’ont pas toutes le même axe de cible ; elles décrivent la trajectoire de conception et ne sont pas un classement apparié unique. La borne finale a été optimisée avec les96épaisseurs à partir d’un96sur saphir.")
    p("Le design retenu a une épaisseur minimale de7,745801nm, maximale de411,902598nm et cumulée de11070,622976nm. Aucune couche de la recette n’est nulle. L’essaiPGLOBAL conjoint supplémentaire dans un voisinage±50nm de borne haute n’a donné aucun gain significatif justifiant de changer le design.")
    p("Le choix du saphir a été étayé par la robustesse indépendante. La recette et la cible sont ensuite figées pendant les études de dépôt. Aucune optimalité globale n’est démontrée.")

    page("5 La réponse spectrale nominale retenue")
    picture(spectrum, "Figure2  Cible finale et spectre nominal du design96sur saphir", 6.1)
    picture(thickness, "Figure3  Répartition des épaisseurs nominales des96couches pièce", 6.1)
    p("RMSE nominal601points :1,4590693866016182point. L’oracle indépendant est passé aux601longueurs d’onde avec un écartR/T maximal8,55×10⁻¹⁵. Le contrôle dense porte sur33485points au pas maximal0,01nm ; les96colonnes du Jacobien ont été vérifiées. [S4]")
    p("Le monitoring à452nm modifie la manière d’arrêter les dépôts, pas la recette nominale ni cette cible. Les CSV de recette et de cible fournis conservent la pleine précision numérique.")

    page("6 Robustesse intrinsèque du design")
    p("Cette étude de robustesse du design précède la recherche de monitoring. Elle applique des erreurs uniformes relatives±0,5%, absolues±0,5nm et leur cumul aux épaisseurs. Elle ne doit pas être confondue avec le bruit de monitoring0,025/0,05/0,10point des campagnesSTRAT. [S5]")
    hold = {r["candidate"]["id"]: r for r in robustness["holdout"]}
    rows = []
    for key, label in [("relative_0p5pct", "Relatif ±0,5%"), ("absolute_0p5nm", "Absolu ±0,5nm"), ("combined", "Cumul")]:
        silica = hold["reference96"]["scenarios"][key]["target_rmse_pp"]
        sapphire = hold["sub3_start728p629"]["scenarios"][key]["target_rmse_pp"]
        rows.append([label, fmt(silica["mean"]), fmt(silica["p95"]), fmt(sapphire["mean"]), fmt(sapphire["p95"])])
    table(["Scénario", "Silice moyen", "Silice P95", "Saphir moyen", "Saphir P95"], rows, [1.7, 1.3, 1.3, 1.3, 1.3], 9.4)
    p("Chaque scénario de holdout comporte10000tirages. En cumul, leP95 passe de3,428245à3,312922points, soit0,115323point de mieux pour la configuration saphir comparée. L’intervalle apparié sur cette différence est rapporté dans l’étude. Les recettes et les cibles des deux configurations ont été réoptimisées : ce résultat ne mesure pas le seul effet isolé du substrat.")
    h("Ce qui est désormais fixé")
    p("La recette finale96saphir, les601ordonnées de cible et la borne728,849591640315nm constituent la référence pour la suite. La recherche de dépôt peut changer modes et longueurs d’onde dans son périmètre autorisé ; elle ne peut s’attribuer un gain en modifiant silencieusement la cible ou les épaisseurs nominales.")
    p("Le fichierdesign_retenu_96_saphir.json permet de reprendre cette référence sans ressaisir les96valeurs. Les anciennes variantes restent identifiables dans les métadonnées archivées.")

    page("7 Contraintes et séquence de fabrication")
    table(["Contrainte", "Application au plan retenu"], [["Témoin", "Un seul témoin saphir"], ["Préparation", "Deux précoatings sur témoin seul"], ["Insertion de la pièce", "Après la préparation sans reset"], ["Pièce", "96couches dont les dix premières optiques"], ["Plafond Rate", "67couches maximum sur96"], ["Répartition retenue", "70optiques et26Rate soit27,08%Rate"], ["Monochromateur", "430à900nm et fente2nm"], ["Opérations totales", "98couches témoin dont96avec la pièce"]], [2.45, 4.45])
    h("Les quatre premières opérations")
    table(["Opération physique", "Destination", "Matériau", "Épaisseur nm", "Monitoring"], [[1, "Témoin seul", "Nb2O5", "130,211377", "452nm"], [2, "Témoin seul", "SiO2", "183,994694", "452nm"], [3, "Pièce1 et témoin", "Nb2O5", fmt(nominal[2]), "452nm"], [4, "Pièce2 et témoin", "SiO2", fmt(nominal[3]), "452nm"]], [1.15, 1.65, 1.05, 1.65, 1.4], 9.3)
    p("L’histoire du témoin reste continue après l’insertion. Les dépôts antérieurs peuvent aider à interpréter le signal des couches suivantes ; aucune information future ne corrige rétroactivement le propre arrêt d’une couche. La continuité de la longueur d’onde évite ici trois changements initiaux510→440→890→480nm.")
    p("Les26couches en Rate sont dénombrées sur la pièce uniquement. La valeur900nm éventuellement présente dans le JSON pour une coucheRate est un emplacement technique et ne représente pas un monitoring optique à900nm.")

    page("8 Méthode de recherche des stratégies de dépôt")
    p("Les stratégies sont évaluées par la croissance native des98couches physiques du témoin et notées sur les96couches réelles de la pièce. La SEEL est calibrée sur la pièce. Les réalisations de perturbation d’indice sont les mêmes dans la croissance et la notation ainsi qu’entre les plans comparés. [S6]")
    p("Le corridor utilisé est additif sur l’indice réel :±0,005sur400à900nm. Il ne s’agit pas de±0,5% relatif. L’absorption est conservée. Les niveaux de bruit sont0,025,0,05et0,10point ; le tableau des paramètres effectifs est joint pour reproduire la configuration exacte.")
    p("Le contrôle optique emploie les historiques mesurés et les règles natives de la brancheSTRAT/POEM. Un point tournant est une propriété du système optique complet : il ne s’identifie pas automatiquement à un quart d’onde de la couche isolée. Une couche épaisse peut transmettre une histoire utile aux suivantes sans autoriser une correction causale de son propre arrêt.")
    table(["Étape statistique", "Usage"], [["Criblage80par bruit", "Explorer rapidement les longueurs d’onde et blocs"], ["Intermédiaire1000par bruit", "Participer au choix des candidats"], ["Confirmation trois graines", "Éprouver des candidats déjà figés"], ["Holdout trois graines fraîches", "Valider des finalistes figés avant résultats"]], [2.7, 4.2])
    p("Fabien a ramené les budgets nouveaux de confirmation et de holdout à1000tirages par bruit et par graine, en conservant trois graines. Les campagnes antérieures à10000ne sont pas réécrites. Les moyennes de graines différentes ne se classent pas directement.")
    p("Pour les variantes de queue, les66premières couches physiques ont été gardées identiques au bit. Cette condition ne s’applique pas à l’essai452nm : les quatre premiers monitoring changent volontairement. Les invariants après l’opération4 ont alors été vérifiés autrement.")

    page("9 Les recherches depuis les plans soixante et soixante quatre")
    p("Les premiers compromis à un témoin ont montré l’intérêt d’ajouter des blocs optiques dans la queue du filtre, après les64premières couches pièce. Les plans hors plafondRate et les variantes à plusieurs témoins sont conservés comme essais, sans devenir la solution choisie. [S6]")
    main_summary = read(REPORTS / "strat60_vaste_2026-10-17/summary.json")
    main_front = main_summary.get("pareto_holdout", [])
    if not main_front:
        main_front = read(REPORTS / "strat60_vaste_2026-10-17/pareto_holdout.json")
    table(["Branche", "Graines finales", "Résultat utile"], [["Élargissement depuis60", "2026108301à8303\n10000par bruit", "67minimumRMSE\n69maximumoptique confirmé"], ["Comparaison fixe incluant64", "2026108501à8503\n10000par bruit", "67,68,69 battent64\nsur chaque graine"], ["Rateépais et POEM depuis64", "2026109801à9803\n10000par bruit", "69et70 battent64\nsur chaque graine"]], [2.5, 2, 2.4], 9.5)
    p("Dans la comparaison fixe, le64est à1,697801point en moyenne, les67/68/69à1,684940/1,687264/1,695151. Les trois candidats terminent tous les dépôts aux trois bruits et améliorent leur référence appariée sur chaque graine. [S6]")
    p("La branchePOEM qualifie70optiques/26Rate, moyenne1,693671contre64apparié1,697451 sur ses propres graines9801à9803. Ce70devient la configuration préférée par Fabien avant l’essai de longueur d’onde commune. Il reste la base exacte du452nm ; les nouveaux réglages de queue des71/74/75n’y ont pas été injectés.")
    p("La différence entre1,684940et1,693671ne prouve aucune supériorité directe entre ces plans car ces deux valeurs proviennent de tirages différents. Chaque conclusion publiée reste liée à sa référence et à sa cohorte.")

    page("10 Variantes autour du soixante dix et branches arrêtées")
    p("La recherche suivante conserve les70modes optiques existants et leurs neuf blocs de queue indivisibles. Elle explore6282plans initiaux et84fenêtres, affine les longueurs d’onde au pas1nm puis combine les parents sur trois époques. Les combinaisons peuvent dépasser75optiques. [S7]")
    vref = variant_audit["reference70_pooled_statistics"][1]
    table(["Optiques", "RMSE moyen", "P95RMSE", "SEELP95 nm"], [["Référence70", fmt(vref["target_rmse_mean_pp"]), fmt(vref["target_rmse_p95_pp"]), fmt(vref["seel_p95_nm"])]] + [[g["optical_component_count"], fmt(g["mean_objective_pp"]), fmt(g["noise_results"][1]["target_rmse_p95_pp"]), fmt(g["noise_results"][1]["seel_p95_nm"])] for g in sorted(variants["pareto_holdout"], key=lambda r:r["optical_component_count"])], [1.3, 1.8, 1.8, 2])
    p("Ces résultats sont des holdouts indépendants20261011201à11203,1000tirages par bruit et graine. Les71/74/75améliorent leur référence70sur chaque graine sans arrêt aux trois bruits.71est le minimumRMSE de cette branche ;75est son maximum optique qualifié.12366évaluations logiques et3150000trajectoires,3182640demandées réellement avec les reprises.")
    h("Le STRAT classique CERTUS")
    p("Le bras classique adapté utilise leWorkerThread23natif,ELITE6,Rate240et joints480, avec adaptations de préfixe, score pièce etPhaseA sur historiquesRate. Il a été arrêté sur demande humaine après29594évaluations et8configurations achevées sur36. Il n’a pas démontré de victoire appariée au70 ; ni raffinement final ni validation indépendante n’ont été achevés. Il ne doit pas être relancé automatiquement. [S9]")
    h("Les pistes à quatre vingts couches optiques")
    p("Treize plans80optiques/16Rate terminent les80tirages par bruit de leur criblagePOEM. Le meilleur a une moyenne1,754022contre64apparié1,696258, soit un coût de0,057763point. Aucun80n’a été qualifié indépendamment dans ce travail. Ce résultat ne justifie pas de remplacer le plan retenu. [S9]")

    page("11 Essai commun aux quatre premières couches")
    p("Fabien a demandé de conserver une même longueur d’onde pendant les deuxprécoatings du témoin et les deuxpremières couches pièce. L’essai lie uniquement ces quatre valeurs, conserve toutes les épaisseurs nominales et le monitoring pièce3à96, et utilise exactement le70POEM précédent comme référence. [S8]")
    picture(lambda_fig, "Figure4  Criblage de123longueurs d’onde communes sur80tirages par bruit")
    p("Le balayage comprend95valeurs de430à900nm par pas5nm, puis28valeurs nouvelles au pas1nm autour de cinq parents. Les points de la figure montrent un score conditionnel aux dépôts terminés ; la présence d’un arrêt reste un critère distinct. Un minimum de cette figure n’est pas une qualification.")
    table(["Phase", "Budget et issue"], [["Criblage", "123λcommunes uniques et référence80"], ["Intermédiaire", "5candidats et référence1000"], ["Confirmation", "4candidats sur trois graines×1000"], ["Holdout candidat", "Aucun finaliste satisfaisant zéro arrêt"], ["Holdout référence", "3références70 mesurées séparément"]], [2.5, 4.4])
    p("Le calcul s’est terminé après148évaluations et101760trajectoires. L’audit a relu toutes les148archives et les444ensembles par bruit, contrôlé400entrées et346sources, les oracles et la chronologie des finalistes. Le statut du protocole est terminé ; le statut scientifique du452nm demeure confirmation seulement.")

    page("12 Résultat du candidat commun à quatre cent cinquante deux")
    p("Sur les mêmes graines de confirmation20261012101à12103,452nm réduit le RMSE moyen à0,05de1,691515à1,667858point. Le gain apparié est0,023657point. LesP95sont calculés sur les3000tirages regroupés par bruit. [S8]")
    table(["Bruit", "RMSE moyen\n452nm", "Référence70", "P95RMSE\n452nm", "P95SEEL\nnm"], [[fmt(a["noise_pp"],3), fmt(a["target_rmse_mean_pp"]), fmt(r["target_rmse_mean_pp"]), fmt(a["target_rmse_p95_pp"]), fmt(a["seel_p95_nm"])] for a,r in zip(retained["pooled_statistics"],base,strict=True)], [1.05, 1.5, 1.5, 1.5, 1.35], 9.5)
    picture(cdf, "Figure5  Distribution appariée au bruit0,05 sur les trois graines de confirmation")
    p("À0,05, leP95RMSE vaut1,870153contre1,935171 ; laSEELP95vaut0,833795contre0,904171nm. Aucun arrêt n’est observé à0,025ou0,05. Au bruit0,10, un tirage sur3000s’arrête pour452nm, contre aucun pour la référence. Les statistiques à0,10portent donc sur2999dépôts terminés.")
    p("Le gain observé ne franchit pas le critère de qualification complet parce que celui-ci imposait zéro arrêt aux trois bruits. Le choix explicite de Fabien conserve452nm comme solution la plus prometteuse avec cette réserve.")

    page("13 Confirmation par graine et réserve de rendement")
    confirmation_rows = [r for r in lambda_rows if r["stage"] == "CONFIRMATION" and r["common_wavelength_nm"] == 452]
    reference_rows = {r["seed"]:r for r in lambda_rows if r["stage"] == "CONFIRMATION_REFERENCE"}
    table(["Graine", "Moyen452", "Référence70", "Gain", "Arrêts\n0,025 0,05 0,10"], [[str(r["seed"]),fmt(r["mean_objective_pp"]),fmt(reference_rows[r["seed"]]["mean_objective_pp"]),fmt(r["mean_gain_pp"])," / ".join(str(t["runs"]-t["completed"]) for t in r["noise_results"])] for r in confirmation_rows], [1.65,1.3,1.3,1.2,1.45],9.2)
    p("L’unique arrêt apparaît à la graine20261012103, au bruit0,10, pour le tirage d’indice888. Le taux observé à ce bruit est1/3000, environ0,0333%. Trois mille tirages ne permettent pas d’identifier exactement le taux physique d’échec ; le résultat ne doit pas être arrondi à zéro.")
    table(["λcommune", "RMSE moyen à0,05", "Arrêts regroupés aux trois bruits"], [[g["common_wavelength_nm"],fmt(g["pooled_statistics"][1]["target_rmse_mean_pp"])," / ".join(str(t["failures"]) for t in g["pooled_statistics"])] for g in lambda_audit["confirmation_groups"]], [1.4,2.4,3.1])
    p("Les quatre candidats449,452,453et454nm ont chacun un arrêt sur le même indice888au bruit0,10. Aucun n’a donc été envoyé aux graines de validation finale12201à12203. Les seules évaluations sur ces graines sont celles de la référence70 ; elles ne valident pas452nm.")
    p("Le plan retenu n’est ni une qualification expérimentale ni une preuve d’optimum global. Une reprise pourra décider explicitement si ce risque observé est acceptable, s’il faut diagnostiquer la cause du tirage888ou réserver de nouvelles graines pour une validation indépendante. Aucun de ces calculs n’a été lancé par la clôture.")

    page("14 Stratégie exacte conservée pour la reprise")
    p("La stratégie retenue est le plan70POEM d’origine avec452nm appliqué aux opérations physiques1à4. Les épaisseurs et les monitoring suivants restent ceux de ce plan. Le fichierstrategie_retenue_452nm.json est une copie exacte du plan testé, et non une combinaison inédite des gagnants de plusieurs campagnes.")
    h("Couches pièce en Rate")
    p(spans(component_rates) + ".")
    h("Couches pièce optiques")
    p(spans(component_optics) + ".")
    table(["Bloc optique de queue", "λ en nm"], [["65à66",519],["67",465],["70à71",530],["75",535],["77",544],["79",535],["81à82",544],["83à86",535],["89",540]], [4.3,2.6])
    p("La numérotation ci-dessus exclut les deuxprécoatings. Pour programmer le monitoring, employer le tableau98opérations des pages suivantes et le CSV associé. L’épaisseur nominale d’une coucheRate reste spécifiée même si aucuneλde monitoring n’est affectée à cette couche.")
    p("Les épaisseurs du tableau imprimé sont arrondies à six décimales. Les JSON et CSV joints sont les références de pleine précision. Les données de simulation restent exprimées en nm et points de pourcentage.")

    for start in range(0,98,25):
        page(f"15 Plan de dépôt opérations {start+1} à {min(start+25,98)}")
        table(["Opération\ntémoin", "Couche\npièce", "Matériau", "Épaisseur\nnominale nm", "Mode", "λ nm"], [[r["physical_layer"],"Précoat" if r["component_layer"] is None else r["component_layer"],"Nb2O5" if r["material"]=="H800-Nb2O5" else "SiO2",fmt(r["nominal_thickness_nm"]),r["mode"],"-" if r["wavelength_nm"] is None else f"{r['wavelength_nm']:g}"] for r in monitoring[start:start+25]], [.85,.85,1.05,1.75,1.3,1.1],9.4)
        p("Pièce insérée après l’opération2, sans reset du témoin. Les précoatings sont exclus du décompte70optiques/26Rate de la pièce. Valeurλabsente surRate : aucune commande optique indiquée.","Caption")

    for start in range(0,100,25):
        page(f"Annexe A Design initial couches {start+1} à {min(start+25,100)}")
        table(["Couche", "Matériau", "QWOT à550nm", "Épaisseur H800 nm"], [[i+1,"Nb2O5" if initial["material_indices"][i]==0 else "SiO2",fmt(initial["QWOT"][i]),fmt(initial["thickness_nm"][i])] for i in range(start,min(start+25,100))],[.85,1.4,2.15,2.5],9.5)
        p("Ordre substrat vers air. QWOT documentaire conservé ; épaisseur physique obtenue avec les indicesH800 à550nm. Le classeur source original est conservé dans l’archive compacte.","Caption")

    page("Annexe B Reprendre le sujet sans ambiguïté")
    p("Bonne Mère est clos à la demande de Fabien afin de revenir au codage de CERTUS. Aucun processus de calcul ne reste actif et le suivi automatique est suspendu. Le plan à452nm est la décision courante ; le70àquatreλdistinctes reste une référence antérieure.")
    h("Ordre de lecture à une reprise future")
    for text in ["LirePASSATION.md puis le présent dossier etdecision_retenue.json.","Vérifier le manifeste des fichiers conservés avant de les utiliser ; ne pas ressaisir les épaisseurs depuis un tableau arrondi.","Confirmer la recette96saphir, lesprécoatings, le70/26et452nm sur les quatre premières opérations.","Relire le bilan0/0/1sur3000et préciser le critère de rendement souhaité avant tout nouveau calcul.","Si une nouvelle étude est décidée, réserver de nouvelles graines et figer le protocole avant résultats ; ne pas présenter la recherche passée comme un holdout frais.","Pour le codage de CERTUS, repartir du préflight et deCLAUDE.md ; ce dossier ne modifie aucune source de production."]:
        p(text,"List Number")
    h("Ce qui reste à examiner au moment opportun")
    p("La cause de l’arrêt sur le tirage888, une éventuelle validation indépendante de452nm et la qualification physique des indices, du saphir, de la fente et du bruit demeurent ouverts. Le modèleRate n’inclut ni dérive de vitesse ni erreur d’outillage témoin/pièce. La conformité machine et les critères de réceptionT/R restent à fixer avant fabrication.")
    h("Nettoyage sous la limite de un gigaoctet")
    p("Le répertoireCERTUS complet, Bonne Mère comprise, doit rester sous1000000000octets. Les sorties brutes volumineuses des candidats non retenus sont supprimées après conservation des métadonnées, des scripts, du moteur figé, des sources et des résultats finaux. L’archive compacte n’est pas une copie intégrale de tous lesNPZ etpickle historiques.")
    p("Les sixcohortes appariées452nm et référence en confirmation sont conservées sous forme compacte :98épaisseurs physiques une seule fois, scores, SEEL et masques de validité. Cette représentation restitue exactement au bit les96épaisseurs pièce et les66premières du témoin. Les audits datés attestent les vérifications faites avant réduction ; les audits historiques complets ne sont plus tous rejouables sans recalculer les sorties éliminées.")

    page("Annexe C Sources et preuves à conserver")
    citations = [
        ("S1", "Documents et modèle", "sources/BM.docx ; sources/BM.pptx ; sources/data.xlsx ; reports/inputs.json"),
        ("S2", "Cible adoptée", "sources/target_retrouvee.txt ; target/target_historique.csv ; design/retenu_96_saphir/target.csv"),
        ("S3", "Design initial et recherches", "design/initial_document_QWOT_H800.json ; reports/qwot_global_historique_01 ; reports/local_116_sans_sio2_2026-10-17 ; reports/local_96_sans_plancher_2026-10-17"),
        ("S4", "Design96saphir et contrôles", "design/retenu_96_saphir/design.json ; selected_validation.json ; dense_validation.json"),
        ("S5", "Robustesse du design", "reports/robustesse96_2026-10-17/summary.json"),
        ("S6", "Dépôt60 et64", "reports/autonomie_nuit_2026-10-17/final_audit_strat60_20261018.json ; final_audit_poem_20261018_0546.json ; final_audit_paired64_20261018_0632.json"),
        ("S7", "Variantes du70", "reports/autonomie_nuit_2026-10-17/final_audit_variants70_20261018_1111.json ; final_cost_variants70_20261018_1115.json"),
        ("S8", "Quatreλcommunes et décision", "reports/strat70_lambda4_2026-10-18/attempt_01/summary.json ; reports/autonomie_nuit_2026-10-17/final_audit_lambda4_20261018_1126.json ; design/deposition_selection.json"),
        ("S9", "Classique arrêté et pistes80", "reports/strat_classique_rate_elargi_2026-10-17/attempt_04/user_stop_summary_20261018.json ; reports/autonomie_nuit_2026-10-17/existing_80_optical_review_20261018.json"),
    ]
    p("Les chemins ci-dessous sont relatifs à l’ancienne racine d’étude. L’archive archives/Bonne_Mere_sources_et_metadonnees.zip conserve les textes, paramètres et sources. Son index archive_index.json associe chaque chemin d’origine à son contenu exact, avec déduplication. L’outil restaurer_archive.py restitue les chemins sur demande. Le plan retenu et ses preuves compactes sont directement dans dossier_final/donnees. L’inventaire de réduction indique les sorties brutes éliminées.")
    for code,title,path in citations:
        h(f"{code} {title}")
        paragraph=p(path)
        for run in paragraph.runs:run.font.size=Pt(9)
        source_list.append(dict(id=code,title=title,original_paths=path))
    save(DATA / "sources_documentaires.json", source_list)
    doc.core_properties.title = "Bonne Mère du design initial à la stratégie de dépôt retenue"
    doc.core_properties.subject = "Clôture et reprise différée du choix70optiques26Rate à452nm"
    doc.core_properties.author = "Dossier préparé pour Fabien"
    for border in doc.styles._element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    for border in doc._element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    output = FINAL / "Bonne_Mere_Dossier_Complet.docx"
    doc.save(output)
    manifest = {p.relative_to(FINAL).as_posix():dict(bytes=p.stat().st_size,sha256=sha(p)) for p in DATA.iterdir() if p.is_file()}
    save(FINAL / "manifest_donnees.json", manifest)
    print(json.dumps(dict(docx=str(output),component_optical_count=70,rate_count=26,monitoring_rows=len(monitoring),
                          source_initial_layers=100,retained_design_layers=96,retained_lambda_nm=452,
                          retained_raw_arrays_bytes=(DATA/"preuves_appariees_452nm_3graines.npz").stat().st_size,
                          figures=len(list(FIG.glob("*.png"))),no_optical_calculation=True),ensure_ascii=True))


if __name__ == "__main__":
    main()
